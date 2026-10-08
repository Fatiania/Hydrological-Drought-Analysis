#!/usr/bin/env python3
"""
STEP 1 - DAILY STREAMFLOW QUALITY SCREENING AND FLOW-REGIME CLASSIFICATION

Purpose
-------
This script prepares a defensible station set for hydrological-drought analysis.
It performs daily-QC checks, annual completeness screening, station screening,
flow-regime classification, FDC threshold diagnostics, and publication-quality
figures/tables.

Core project rules
------------------
1) Daily discharge QC
   - Sentinel values -999 and -9999 -> NaN.
   - Negative discharge -> NaN.
   - Duplicate dates: if all valid duplicate values agree, keep the value;
     if valid duplicate values conflict, set that date to NaN.
   - A complete daily calendar is inserted.
   - Consecutive zero runs (>=2 days) are retained.
   - An isolated Q=0 with both immediate neighbours valid is replaced by the
     arithmetic mean of the two neighbours.
   - An isolated Q=0 with either immediate neighbour missing is set to NaN.
   - Missing intervals are NOT interpolated.

2) Annual completeness
   - A calendar year is usable only when it has <=15 missing daily values.
     Therefore a common year needs >=350 valid days and a leap year >=351.
   - This follows the completeness logic used by Fleig et al. (2006), while
     this project does not interpolate the missing days.

3) Station inclusion
   - Retain a station only when it has >=15 usable calendar years.

4) Flow-regime classification
   - Use only observations belonging to usable years.
   - Zero-flow percentage = 100 * N(Q=0) / N(valid daily observations).
   - Classification threshold = 0.5479452055%.
   - Perennial:    zero-flow percentage < threshold.
   - Intermittent: zero-flow percentage >= threshold.

5) Drought-threshold diagnostics and final thresholds
   - Perennial diagnostics: Q95, Q90, Q85; FINAL drought analysis uses Q90.
   - Intermittent diagnostics: Q70, Q60, Q50, Q40, Q30, Q20; FINAL drought analysis uses Q50.

6) Event-based threshold sensitivity
   - Every candidate threshold is applied independently to the daily series.
   - A drought day is Q < Q_threshold.
   - Missing values are never interpolated and always break an event.
   - Unusable years are masked to NaN and therefore also break an event.
   - Events touching a missing/unusable interval or the record boundary are flagged as censored.
   - Onset, recovery, duration, Qmin, normalized Qmin, deficit, zero-flow contribution,
     event frequency, and seasonal timing are calculated for every candidate threshold.

7) Final event-pooling comparison
   - Use Q90 for every retained perennial station and Q50 for every retained intermittent station.
   - Compare Raw (no pooling), IT, IC, MA and SPA.
   - Missing or unusable days always break the analysis and are never pooled across.
   - Final duration and deficit summaries use complete (uncensored) events.
   - Save one five-panel hydrograph per station showing the fixed threshold, event duration,
     and physical deficit for Raw/IT/IC/MA/SPA.
   - Save one all-stations pooling summary figure for perennial stations and one for
     intermittent stations.

Outputs
-------
One consolidated Excel workbook (all result tables as separate sheets, including event catalogues and threshold-sensitivity diagnostics) plus publication-quality figures:
- PNG 600 dpi
- PDF vector
- optional TIFF 600 dpi (enable SAVE_TIFF=True)

Expected station-file columns: Time, Flow
Default input folder: C:\\Drought\\DATA\\MAIN DATA\\STATIONS

Important figure-design note
----------------------------
The figures use publication-quality, color-blind-friendly colors. The previous
zoomed plot around the regime cutoff is deliberately omitted. The numerical
percentage threshold is shown only on the main regime diagnostic figure; no
day-based wording is written on the plots.
"""

from __future__ import annotations

import argparse
import calendar
import gc
import math
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import ColorScaleRule

warnings.filterwarnings("ignore", category=RuntimeWarning)

# =============================================================================
# USER SETTINGS
# =============================================================================
DEFAULT_INPUT = Path(r"C:\Drought\DATA\MAIN DATA\STATIONS")
DEFAULT_OUTPUT = Path(r"C:\Drought\DATA\MAIN DATA\STEP1_QUALITY_SCREENING_COLOR")

MAX_MISSING_DAYS_PER_USABLE_YEAR = 15
MIN_USABLE_YEARS = 15
REGIME_ZERO_PCT_CUTOFF = 100.0 * 2.0 / 365.0  # 0.5479452055%; plots show only percentage

OUTPUT_DPI = 600
SAVE_PDF = True
SAVE_TIFF = False
EXCEL_FILENAME = "Step1_Drought_Quality_Screening_All_Results.xlsx"

# Final drought-analysis thresholds requested for the pooling comparison.
FINAL_THRESHOLD_PERENNIAL = "Q90"
FINAL_THRESHOLD_INTERMITTENT = "Q50"

# Pooling-method settings. These are analyst-controlled settings, not universal constants.
# They match the examples discussed in the accompanying presentation.
IT_CRITICAL_TIME_DAYS = 5
IC_CRITICAL_TIME_DAYS = 5
IC_CRITICAL_VOLUME_RATIO = 0.10
MA_WINDOW_DAYS = 10

POOLING_METHOD_ORDER = ["Raw", "IT", "IC", "MA", "SPA"]
POOLING_METHOD_LABELS = {
    "Raw": "Raw (no pooling)",
    "IT": "IT",
    "IC": "IC",
    "MA": "MA",
    "SPA": "SPA",
}
POOLING_COLORS = {
    "Raw": "#7F7F7F",
    "IT": "#0072B2",
    "IC": "#009E73",
    "MA": "#E69F00",
    "SPA": "#CC79A7",
}

# Publication style
plt.rcParams.update({
    "font.family": "serif",
    "font.serif": ["Times New Roman", "DejaVu Serif"],
    "font.size": 10,
    "axes.titlesize": 11,
    "axes.labelsize": 10,
    "xtick.labelsize": 8,
    "ytick.labelsize": 8,
    "legend.fontsize": 10,
    "figure.dpi": 150,
    "savefig.dpi": OUTPUT_DPI,
    "axes.linewidth": 0.8,
})

# Color-blind-friendly publication palette
COLOR_RETAINED = "#009E73"       # green
COLOR_EXCLUDED = "#A6A6A6"       # neutral gray
COLOR_PERENNIAL = "#0072B2"      # blue
COLOR_INTERMITTENT = "#D55E00"   # vermillion
COLOR_Q95 = "#56B4E9"            # light blue
COLOR_Q90 = "#0072B2"            # blue
COLOR_Q85 = "#E69F00"            # orange
COLOR_Q70 = "#CC79A7"            # purple/pink
COLOR_Q60 = "#D55E00"            # vermillion
COLOR_Q50 = "#E69F00"            # orange
COLOR_Q40 = "#009E73"            # green
COLOR_Q30 = "#56B4E9"            # sky blue
COLOR_Q20 = "#0072B2"            # blue
COLOR_USABLE = "#2CA25F"         # green
COLOR_NOT_USABLE = "#FEE08B"     # warm yellow


# =============================================================================
# UTILITIES
# =============================================================================
def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Quality-controlled streamflow screening for drought analysis."
    )
    p.add_argument("--input", type=Path, default=DEFAULT_INPUT,
                   help="Folder containing station CSV files.")
    p.add_argument("--output", type=Path, default=DEFAULT_OUTPUT,
                   help="Output folder.")
    p.add_argument("--metadata", type=Path, default=None,
                   help="Optional Catchments_info.csv.")
    return p.parse_args()


def write_excel_workbook(tables: list[tuple[str, pd.DataFrame, str]], path: Path) -> None:
    """Write all Step 1 tables to one Excel workbook using fast write-only mode.

    The event catalogue can contain tens of thousands of rows, so write-only
    mode is used to keep memory use and execution time reasonable.
    """
    from openpyxl.cell import WriteOnlyCell

    path.parent.mkdir(parents=True, exist_ok=True)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Times New Roman", size=10, bold=True, color="FFFFFF")
    title_font = Font(name="Times New Roman", size=14, bold=True, color="1F4E78")
    body_font = Font(name="Times New Roman", size=10, color="000000")

    wb = Workbook(write_only=True)

    # README / index sheet
    ws0 = wb.create_sheet("README")
    ws0.sheet_view.showGridLines = False
    ws0.freeze_panes = "A6"
    ws0.column_dimensions["A"].width = 24
    ws0.column_dimensions["B"].width = 10
    ws0.column_dimensions["C"].width = 90

    c = WriteOnlyCell(ws0, value="STEP 1 - Streamflow quality screening and regime classification")
    c.font = title_font
    ws0.append([c])
    ws0.append([])
    c = WriteOnlyCell(ws0, value="All Step 1 result tables, event catalogues, and threshold-sensitivity diagnostics are consolidated in this workbook.")
    c.font = body_font
    ws0.append([c])
    ws0.append([])

    hdr = []
    for value in ["Sheet", "Rows", "Description"]:
        c = WriteOnlyCell(ws0, value=value)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")
        hdr.append(c)
    ws0.append(hdr)
    for sheet_name, df, description in tables:
        ws0.append([sheet_name, len(df), description])

    for sheet_name, df, description in tables:
        ws = wb.create_sheet(sheet_name)
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"

        if len(df.columns) == 0:
            h = WriteOnlyCell(ws, value="Status")
            h.fill = header_fill
            h.font = header_font
            ws.append([h])
            ws.append(["No records were generated for this table in the current analysis."])
            ws.column_dimensions["A"].width = 72
            continue

        # Compute readable widths before writing rows.
        for j, col in enumerate(df.columns, start=1):
            values = [str(col)]
            if len(df) > 0:
                values.extend(df.iloc[:250, j-1].dropna().astype(str).tolist())
            max_len = max((len(v) for v in values), default=10)
            width = min(max(max_len + 2, 10), 36)
            if any(k in str(col) for k in ["Rule", "Historical_Period", "Issue", "Reason"]):
                width = min(max(width, 24), 50)
            ws.column_dimensions[get_column_letter(j)].width = width

        header_cells = []
        for col in df.columns:
            cell = WriteOnlyCell(ws, value=str(col))
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            header_cells.append(cell)
        ws.append(header_cells)

        for row in df.itertuples(index=False, name=None):
            values = []
            for val in row:
                if pd.isna(val):
                    values.append(None)
                else:
                    values.append(val.item() if hasattr(val, "item") else val)
            ws.append(values)

        last_col = get_column_letter(len(df.columns))
        last_row = max(2, len(df) + 1)
        ws.auto_filter.ref = f"A1:{last_col}{last_row}"

    wb.save(path)

def save_figure(fig: plt.Figure, png_path: Path) -> None:
    png_path.parent.mkdir(parents=True, exist_ok=True)
    stem = png_path.with_suffix("")
    fig.savefig(png_path, dpi=OUTPUT_DPI, bbox_inches="tight", pad_inches=0.05,
                facecolor="white")
    if SAVE_PDF:
        fig.savefig(stem.with_suffix(".pdf"), bbox_inches="tight", pad_inches=0.05,
                    facecolor="white")
    if SAVE_TIFF:
        fig.savefig(stem.with_suffix(".tiff"), dpi=OUTPUT_DPI,
                    bbox_inches="tight", pad_inches=0.05, facecolor="white",
                    pil_kwargs={"compression": "tiff_lzw"})


def run_segments(mask: pd.Series) -> list[tuple[int, int]]:
    arr = mask.fillna(False).to_numpy(dtype=bool)
    if arr.size == 0:
        return []
    starts = np.where(arr & np.r_[True, ~arr[:-1]])[0]
    ends = np.where(arr & np.r_[~arr[1:], True])[0]
    return [(int(s), int(e)) for s, e in zip(starts, ends)]


def fdc_q(flow: pd.Series, exceedance_percent: float) -> float:
    vals = pd.to_numeric(flow, errors="coerce").dropna().to_numpy(dtype=float)
    if vals.size == 0:
        return np.nan
    # Qp = discharge equalled/exceeded p% of the time
    return float(np.nanpercentile(vals, 100.0 - float(exceedance_percent)))

def detect_drought_events(
    daily: pd.DataFrame,
    threshold: float,
    station: str,
    regime: str,
    threshold_label: str,
    exceedance_percent: float,
) -> list[dict]:
    """Detect uninterrupted Q < threshold event segments on a complete calendar.

    Missing values are NOT interpolated. A missing day breaks an event.
    An event is left/right censored when it touches a missing/unusable day or
    the beginning/end of the available series. Recovery is the first subsequent
    valid day with Q >= threshold, consistent with defining drought as Q < threshold.
    """
    if daily.empty or not np.isfinite(threshold) or threshold <= 0:
        return []

    d = daily[["Time", "Flow"]].copy().reset_index(drop=True)
    flow = pd.to_numeric(d["Flow"], errors="coerce")
    valid = flow.notna()
    below = valid & (flow < float(threshold))
    segments = run_segments(below)

    rows: list[dict] = []
    for event_id, (s, e) in enumerate(segments, start=1):
        q_event = flow.iloc[s:e+1].astype(float)
        onset = pd.Timestamp(d.loc[s, "Time"])
        last_drought = pd.Timestamp(d.loc[e, "Time"])

        if s == 0:
            left_censored = True
            left_reason = "Start_of_series"
        elif pd.isna(flow.iloc[s-1]):
            left_censored = True
            left_reason = "Missing_or_unusable_day_before"
        else:
            left_censored = False
            left_reason = ""

        if e == len(d) - 1:
            right_censored = True
            right_reason = "End_of_series"
            recovery_date = pd.NaT
            recovery_q = np.nan
        elif pd.isna(flow.iloc[e+1]):
            right_censored = True
            right_reason = "Missing_or_unusable_day_after"
            recovery_date = pd.NaT
            recovery_q = np.nan
        else:
            right_censored = False
            right_reason = ""
            recovery_date = pd.Timestamp(d.loc[e+1, "Time"])
            recovery_q = float(flow.iloc[e+1])

        qmin = float(q_event.min()) if len(q_event) else np.nan
        duration = int(e - s + 1)
        deficit_sum = float((float(threshold) - q_event).sum()) if len(q_event) else np.nan
        zero_days = int(np.isclose(q_event.to_numpy(dtype=float), 0.0, atol=0.0).sum())
        qmin_ratio = qmin / float(threshold) if threshold > 0 and np.isfinite(qmin) else np.nan
        min_deficit_ratio = 1.0 - qmin_ratio if np.isfinite(qmin_ratio) else np.nan

        rows.append({
            "Station": station,
            "Flow_Regime": regime,
            "Threshold_Label": threshold_label,
            "Threshold_Exceedance_pct": float(exceedance_percent),
            "Threshold_m3s": float(threshold),
            "Event_ID": event_id,
            "Onset_Date": onset.date().isoformat(),
            "Last_Drought_Day": last_drought.date().isoformat(),
            "Recovery_Date": recovery_date.date().isoformat() if pd.notna(recovery_date) else None,
            "Recovery_Q_m3s": recovery_q,
            "Duration_days": duration,
            "Qmin_m3s": qmin,
            "Qmin_to_Qthreshold": qmin_ratio,
            "MinFlow_Deficit_Ratio": min_deficit_ratio,
            "Deficit_Sum_m3s_day": deficit_sum,
            "Mean_Daily_Deficit_m3s": deficit_sum / duration if duration else np.nan,
            "Zero_Flow_Days": zero_days,
            "Zero_Flow_Share_of_Event_pct": 100.0 * zero_days / duration if duration else np.nan,
            "Onset_Month": int(onset.month),
            "Onset_DayOfYear": int(onset.dayofyear),
            "Recovery_Month": int(recovery_date.month) if pd.notna(recovery_date) else np.nan,
            "Recovery_DayOfYear": int(recovery_date.dayofyear) if pd.notna(recovery_date) else np.nan,
            "Left_Censored": "Yes" if left_censored else "No",
            "Right_Censored": "Yes" if right_censored else "No",
            "Censored": "Yes" if (left_censored or right_censored) else "No",
            "Left_Censor_Reason": left_reason,
            "Right_Censor_Reason": right_reason,
        })
    return rows


def threshold_event_summary(
    daily: pd.DataFrame,
    threshold: float,
    station: str,
    regime: str,
    threshold_label: str,
    exceedance_percent: float,
    usable_years: int,
) -> tuple[dict, list[dict]]:
    """Return station-level sensitivity metrics and its event catalogue."""
    valid = pd.to_numeric(daily["Flow"], errors="coerce").notna() if not daily.empty else pd.Series(dtype=bool)
    n_valid = int(valid.sum()) if len(valid) else 0

    base = {
        "Station": station,
        "Flow_Regime": regime,
        "Threshold_Label": threshold_label,
        "Threshold_Exceedance_pct": float(exceedance_percent),
        "Threshold_m3s": threshold,
        "Usable_Years": int(usable_years),
        "Usable_Valid_Days": n_valid,
    }

    if not np.isfinite(threshold) or threshold <= 0 or n_valid == 0:
        base.update({
            "Threshold_Status": "Nonpositive_or_unavailable",
            "Observed_Drought_Days": 0,
            "Drought_Days_per_Usable_Year": 0.0 if usable_years else np.nan,
            "Drought_Days_pct_of_Valid": 0.0 if n_valid else np.nan,
            "Observed_Event_Segments": 0,
            "Observed_Event_Segments_per_Year": 0.0 if usable_years else np.nan,
            "Complete_Events": 0,
            "Complete_Events_per_Year": 0.0 if usable_years else np.nan,
            "Censored_Events": 0,
            "Censored_Event_Share_pct": np.nan,
            "Median_Duration_days_Complete": np.nan,
            "Mean_Duration_days_Complete": np.nan,
            "Max_Duration_days_Complete": np.nan,
            "Median_Qmin_m3s_Complete": np.nan,
            "Median_Qmin_to_Qthreshold_Complete": np.nan,
            "Median_MinFlow_Deficit_Ratio_Complete": np.nan,
            "Median_Deficit_Sum_m3s_day_Complete": np.nan,
            "Zero_Flow_Days_within_Drought": 0,
            "Zero_Flow_Share_of_Drought_Days_pct": np.nan,
        })
        return base, []

    flow = pd.to_numeric(daily["Flow"], errors="coerce")
    drought_day_mask = flow.notna() & (flow < float(threshold))
    drought_days = int(drought_day_mask.sum())
    zero_within = int((drought_day_mask & np.isclose(flow, 0.0, atol=0.0)).sum())

    events = detect_drought_events(
        daily, threshold, station, regime, threshold_label, exceedance_percent
    )
    ev = pd.DataFrame(events)
    n_events = len(ev)
    if not ev.empty:
        complete = ev.loc[ev["Censored"] == "No"].copy()
        n_complete = len(complete)
        n_censored = int(n_events - n_complete)
    else:
        complete = pd.DataFrame()
        n_complete = 0
        n_censored = 0

    def med(col):
        return float(complete[col].median()) if n_complete and col in complete else np.nan

    def mean(col):
        return float(complete[col].mean()) if n_complete and col in complete else np.nan

    def maxv(col):
        return float(complete[col].max()) if n_complete and col in complete else np.nan

    base.update({
        "Threshold_Status": "Evaluated",
        "Observed_Drought_Days": drought_days,
        "Drought_Days_per_Usable_Year": drought_days / usable_years if usable_years else np.nan,
        "Drought_Days_pct_of_Valid": 100.0 * drought_days / n_valid if n_valid else np.nan,
        "Observed_Event_Segments": n_events,
        "Observed_Event_Segments_per_Year": n_events / usable_years if usable_years else np.nan,
        "Complete_Events": n_complete,
        "Complete_Events_per_Year": n_complete / usable_years if usable_years else np.nan,
        "Censored_Events": n_censored,
        "Censored_Event_Share_pct": 100.0 * n_censored / n_events if n_events else np.nan,
        "Median_Duration_days_Complete": med("Duration_days"),
        "Mean_Duration_days_Complete": mean("Duration_days"),
        "Max_Duration_days_Complete": maxv("Duration_days"),
        "Median_Qmin_m3s_Complete": med("Qmin_m3s"),
        "Median_Qmin_to_Qthreshold_Complete": med("Qmin_to_Qthreshold"),
        "Median_MinFlow_Deficit_Ratio_Complete": med("MinFlow_Deficit_Ratio"),
        "Median_Deficit_Sum_m3s_day_Complete": med("Deficit_Sum_m3s_day"),
        "Zero_Flow_Days_within_Drought": zero_within,
        "Zero_Flow_Share_of_Drought_Days_pct": 100.0 * zero_within / drought_days if drought_days else np.nan,
    })
    return base, events


def aggregate_regime_sensitivity(station_sens: pd.DataFrame) -> pd.DataFrame:
    """Aggregate station-level threshold sensitivity using median and IQR."""
    if station_sens.empty:
        return pd.DataFrame()

    metrics = [
        "Drought_Days_per_Usable_Year",
        "Complete_Events_per_Year",
        "Median_Duration_days_Complete",
        "Median_MinFlow_Deficit_Ratio_Complete",
        "Zero_Flow_Share_of_Drought_Days_pct",
        "Censored_Event_Share_pct",
    ]
    rows = []
    for (regime, label, exc), g in station_sens.groupby(
        ["Flow_Regime", "Threshold_Label", "Threshold_Exceedance_pct"], dropna=False
    ):
        row = {
            "Flow_Regime": regime,
            "Threshold_Label": label,
            "Threshold_Exceedance_pct": exc,
            "Stations_Total": int(g["Station"].nunique()),
            "Stations_with_Positive_Threshold": int((pd.to_numeric(g["Threshold_m3s"], errors="coerce") > 0).sum()),
        }
        for m in metrics:
            vals = pd.to_numeric(g[m], errors="coerce").dropna()
            row[f"{m}_Median"] = float(vals.median()) if len(vals) else np.nan
            row[f"{m}_Q25"] = float(vals.quantile(0.25)) if len(vals) else np.nan
            row[f"{m}_Q75"] = float(vals.quantile(0.75)) if len(vals) else np.nan
        rows.append(row)
    return pd.DataFrame(rows)


def monthly_event_timing(event_df: pd.DataFrame) -> pd.DataFrame:
    """Percentage distribution of complete-event onset and recovery months."""
    if event_df.empty:
        return pd.DataFrame()
    e = event_df.loc[event_df["Censored"] == "No"].copy()
    if e.empty:
        return pd.DataFrame()
    rows = []
    for (regime, label, exc), g in e.groupby(
        ["Flow_Regime", "Threshold_Label", "Threshold_Exceedance_pct"], dropna=False
    ):
        for timing, col in [("Onset", "Onset_Month"), ("Recovery", "Recovery_Month")]:
            vals = pd.to_numeric(g[col], errors="coerce").dropna().astype(int)
            total = len(vals)
            for month in range(1, 13):
                n = int((vals == month).sum())
                rows.append({
                    "Flow_Regime": regime,
                    "Threshold_Label": label,
                    "Threshold_Exceedance_pct": exc,
                    "Timing": timing,
                    "Month": month,
                    "Event_Count": n,
                    "Event_Percentage": 100.0 * n / total if total else np.nan,
                })
    return pd.DataFrame(rows)


# =============================================================================
# FINAL DROUGHT ANALYSIS: FIXED THRESHOLDS + POOLING APPROACHES
# =============================================================================
def _original_deficit(flow: pd.Series, threshold: float) -> float:
    """Positive streamflow shortfall below the fixed threshold, in m3/s-day."""
    q = pd.to_numeric(flow, errors="coerce").to_numpy(dtype=float)
    q = q[np.isfinite(q)]
    if q.size == 0:
        return np.nan
    return float(np.maximum(float(threshold) - q, 0.0).sum())


def _raw_threshold_intervals(daily: pd.DataFrame, threshold: float) -> list[dict]:
    """Raw uninterrupted Q < threshold intervals. Missing days always break events."""
    q = pd.to_numeric(daily["Flow"], errors="coerce")
    mask = q.notna() & (q < float(threshold))
    rows = []
    for s, e in run_segments(mask):
        deficit = _original_deficit(q.iloc[s:e+1], threshold)
        rows.append({
            "start": int(s),
            "end": int(e),
            "raw_count": 1,
            "method_severity": deficit,
        })
    return rows


def _pool_it_intervals(daily: pd.DataFrame, threshold: float,
                       critical_days: int) -> list[dict]:
    """Inter-event Time criterion: merge when the valid inter-event gap <= tc."""
    q = pd.to_numeric(daily["Flow"], errors="coerce")
    raw = _raw_threshold_intervals(daily, threshold)
    if not raw:
        return []

    pooled = []
    current = dict(raw[0])
    for nxt0 in raw[1:]:
        nxt = dict(nxt0)
        gap_start = current["end"] + 1
        gap_end = nxt["start"] - 1
        gap_days = max(0, gap_end - gap_start + 1)
        gap = q.iloc[gap_start:gap_end+1] if gap_days else pd.Series(dtype=float)
        valid_gap = bool(gap_days > 0 and gap.notna().all())

        if valid_gap and gap_days <= int(critical_days):
            current["end"] = nxt["end"]
            current["raw_count"] += nxt["raw_count"]
            # IT merges by time only. The physical deficit is the sum of positive
            # shortfalls; above-threshold inter-event days add no deficit.
            current["method_severity"] = float(current["method_severity"] + nxt["method_severity"])
        else:
            pooled.append(current)
            current = nxt
    pooled.append(current)
    return pooled


def _pool_ic_intervals(daily: pd.DataFrame, threshold: float,
                       critical_days: int, critical_ratio: float) -> list[dict]:
    """Inter-event Time and Volume criterion following the slide definition.

    Merge neighbouring raw droughts when both:
        ti <= tc
        Vi / S1 <= pc
    where Vi is excess flow above the threshold during the inter-event period and
    S1 is the current event deficit. After merging, method-specific net severity is
    S1 + S2 - Vi. Missing/unusable days never permit pooling.
    """
    q = pd.to_numeric(daily["Flow"], errors="coerce")
    raw = _raw_threshold_intervals(daily, threshold)
    if not raw:
        return []

    pooled = []
    current = dict(raw[0])
    for nxt0 in raw[1:]:
        nxt = dict(nxt0)
        gap_start = current["end"] + 1
        gap_end = nxt["start"] - 1
        gap_days = max(0, gap_end - gap_start + 1)
        gap = q.iloc[gap_start:gap_end+1] if gap_days else pd.Series(dtype=float)
        valid_gap = bool(gap_days > 0 and gap.notna().all())

        if valid_gap:
            excess = float(np.maximum(gap.to_numpy(dtype=float) - float(threshold), 0.0).sum())
        else:
            excess = np.nan

        s1 = float(current.get("method_severity", np.nan))
        ratio = excess / s1 if valid_gap and np.isfinite(s1) and s1 > 0 else np.inf

        if valid_gap and gap_days <= int(critical_days) and ratio <= float(critical_ratio):
            s2 = float(nxt.get("method_severity", 0.0))
            current["end"] = nxt["end"]
            current["raw_count"] += nxt["raw_count"]
            current["method_severity"] = max(0.0, s1 + s2 - excess)
        else:
            pooled.append(current)
            current = nxt
    pooled.append(current)
    return pooled


def _moving_average_intervals(daily: pd.DataFrame, threshold: float,
                              window_days: int) -> list[dict]:
    """Moving Average approach applied independently within each valid block."""
    q = pd.to_numeric(daily["Flow"], errors="coerce")
    ma = pd.Series(np.nan, index=daily.index, dtype=float)
    valid = q.notna()

    for s, e in run_segments(valid):
        block = q.iloc[s:e+1].astype(float)
        smoothed = block.rolling(
            window=int(window_days), center=True, min_periods=int(window_days)
        ).mean()
        ma.iloc[s:e+1] = smoothed.to_numpy(dtype=float)

    below = ma.notna() & (ma < float(threshold))
    rows = []
    for s, e in run_segments(below):
        sev = float((float(threshold) - ma.iloc[s:e+1]).clip(lower=0).sum())
        rows.append({
            "start": int(s),
            "end": int(e),
            "raw_count": np.nan,
            "method_severity": sev,
        })
    return rows


def _spa_intervals(daily: pd.DataFrame, threshold: float) -> list[dict]:
    """Sequent Peak Algorithm using accumulated deficit w(t)."""
    q = pd.to_numeric(daily["Flow"], errors="coerce")
    rows = []

    for bs, be in run_segments(q.notna()):
        w = 0.0
        event_start = None
        peak_w = 0.0
        for i in range(bs, be + 1):
            qi = float(q.iloc[i])
            w = max(0.0, w + float(threshold) - qi)

            if event_start is None and w > 0:
                event_start = i
                peak_w = w
            elif event_start is not None:
                peak_w = max(peak_w, w)
                if w <= 0:
                    # The current day is the recovery day; the drought interval
                    # terminates on the preceding day.
                    end = i - 1
                    if end >= event_start:
                        rows.append({
                            "start": int(event_start),
                            "end": int(end),
                            "raw_count": np.nan,
                            "method_severity": float(peak_w),
                        })
                    event_start = None
                    peak_w = 0.0

        if event_start is not None:
            rows.append({
                "start": int(event_start),
                "end": int(be),
                "raw_count": np.nan,
                "method_severity": float(peak_w),
            })
    return rows


def _intervals_to_event_rows(daily: pd.DataFrame, intervals: list[dict],
                             station: str, regime: str, threshold_label: str,
                             threshold: float, method: str,
                             parameter_text: str) -> list[dict]:
    """Convert final method intervals to a common, comparable event catalogue."""
    if not intervals:
        return []

    q = pd.to_numeric(daily["Flow"], errors="coerce")
    rows = []
    for event_id, iv in enumerate(intervals, start=1):
        s = int(iv["start"])
        e = int(iv["end"])
        if s < 0 or e < s or e >= len(daily):
            continue

        span_q = q.iloc[s:e+1]
        onset = pd.Timestamp(daily.iloc[s]["Time"])
        last_day = pd.Timestamp(daily.iloc[e]["Time"])
        duration = int(e - s + 1)

        left_censored = bool(s == 0 or pd.isna(q.iloc[s-1]))
        right_censored = bool(e == len(daily)-1 or pd.isna(q.iloc[e+1]))
        if not right_censored:
            termination = pd.Timestamp(daily.iloc[e+1]["Time"])
        else:
            termination = pd.NaT

        physical_deficit = _original_deficit(span_q, threshold)
        drought_days = int((span_q.notna() & (span_q < float(threshold))).sum())
        zero_days = int((span_q.notna() & np.isclose(span_q, 0.0, atol=0.0)).sum())
        min_q = float(span_q.min()) if span_q.notna().any() else np.nan
        deficit_volume = physical_deficit * 86400.0 if np.isfinite(physical_deficit) else np.nan

        rows.append({
            "Station": station,
            "Flow_Regime": regime,
            "Final_Threshold_Label": threshold_label,
            "Final_Threshold_m3s": float(threshold),
            "Pooling_Method": method,
            "Pooling_Parameters": parameter_text,
            "Event_ID": event_id,
            "Onset_Date": onset.date().isoformat(),
            "Last_Event_Day": last_day.date().isoformat(),
            "Termination_Date": termination.date().isoformat() if pd.notna(termination) else None,
            "Duration_days": duration,
            "Drought_Days_Within_Event": drought_days,
            "Interevent_or_Recovery_Days_Within_Event": int(duration - drought_days),
            "Minimum_Flow_m3s": min_q,
            "Deficit_Sum_m3s_day": physical_deficit,
            "Deficit_Volume_m3": deficit_volume,
            "Mean_Deficit_over_Event_m3s": physical_deficit / duration if duration and np.isfinite(physical_deficit) else np.nan,
            "Method_Specific_Severity_m3s_day": iv.get("method_severity", np.nan),
            "Raw_Constituent_Events": iv.get("raw_count", np.nan),
            "Zero_Flow_Days": zero_days,
            "Left_Censored": "Yes" if left_censored else "No",
            "Right_Censored": "Yes" if right_censored else "No",
            "Censored": "Yes" if (left_censored or right_censored) else "No",
        })
    return rows


def final_pooling_analysis(daily: pd.DataFrame, threshold: float, station: str,
                           regime: str, threshold_label: str,
                           usable_years: int) -> tuple[list[dict], list[dict]]:
    """Run Raw, IT, IC, MA and SPA using one fixed threshold for the station."""
    if daily.empty or not np.isfinite(threshold) or threshold <= 0:
        return [], []

    method_intervals = {
        "Raw": _raw_threshold_intervals(daily, threshold),
        "IT": _pool_it_intervals(daily, threshold, IT_CRITICAL_TIME_DAYS),
        "IC": _pool_ic_intervals(
            daily, threshold, IC_CRITICAL_TIME_DAYS, IC_CRITICAL_VOLUME_RATIO
        ),
        "MA": _moving_average_intervals(daily, threshold, MA_WINDOW_DAYS),
        "SPA": _spa_intervals(daily, threshold),
    }
    parameter_text = {
        "Raw": "No pooling",
        "IT": f"tc={IT_CRITICAL_TIME_DAYS} days",
        "IC": f"tc={IC_CRITICAL_TIME_DAYS} days; pc={IC_CRITICAL_VOLUME_RATIO:.2f}",
        "MA": f"centered {MA_WINDOW_DAYS}-day moving average",
        "SPA": "w(t)=max[0, w(t-1)+Qthr-Q(t)]",
    }

    event_rows: list[dict] = []
    summary_rows: list[dict] = []

    for method in POOLING_METHOD_ORDER:
        rows = _intervals_to_event_rows(
            daily=daily,
            intervals=method_intervals[method],
            station=station,
            regime=regime,
            threshold_label=threshold_label,
            threshold=threshold,
            method=method,
            parameter_text=parameter_text[method],
        )
        event_rows.extend(rows)
        ev = pd.DataFrame(rows)
        if ev.empty:
            complete = pd.DataFrame()
            n_all = 0
            n_complete = 0
            n_censored = 0
        else:
            complete = ev.loc[ev["Censored"] == "No"].copy()
            n_all = len(ev)
            n_complete = len(complete)
            n_censored = n_all - n_complete

        def stat(col: str, func: str):
            if complete.empty or col not in complete:
                return np.nan
            vals = pd.to_numeric(complete[col], errors="coerce").dropna()
            if vals.empty:
                return np.nan
            if func == "median":
                return float(vals.median())
            if func == "mean":
                return float(vals.mean())
            if func == "max":
                return float(vals.max())
            raise ValueError(func)

        summary_rows.append({
            "Station": station,
            "Flow_Regime": regime,
            "Final_Threshold_Label": threshold_label,
            "Final_Threshold_m3s": float(threshold),
            "Pooling_Method": method,
            "Pooling_Parameters": parameter_text[method],
            "Usable_Years": int(usable_years),
            "All_Events": int(n_all),
            "Complete_Events": int(n_complete),
            "Censored_Events": int(n_censored),
            "Complete_Events_per_Usable_Year": n_complete / usable_years if usable_years else np.nan,
            "Median_Duration_days": stat("Duration_days", "median"),
            "Mean_Duration_days": stat("Duration_days", "mean"),
            "Max_Duration_days": stat("Duration_days", "max"),
            "Median_Deficit_m3s_day": stat("Deficit_Sum_m3s_day", "median"),
            "Mean_Deficit_m3s_day": stat("Deficit_Sum_m3s_day", "mean"),
            "Max_Deficit_m3s_day": stat("Deficit_Sum_m3s_day", "max"),
            "Median_Deficit_Volume_m3": stat("Deficit_Volume_m3", "median"),
            "Mean_Deficit_Volume_m3": stat("Deficit_Volume_m3", "mean"),
            "Max_Deficit_Volume_m3": stat("Deficit_Volume_m3", "max"),
        })

    return summary_rows, event_rows


def plot_pooling_comparison_by_station(pooling_summary: pd.DataFrame,
                                       pooling_events: pd.DataFrame,
                                       fig_dir: Path) -> None:
    """One station-specific figure comparing Raw, IT, IC, MA and SPA."""
    if pooling_summary.empty:
        return

    base_dir = fig_dir / "13_Final_Pooling_Comparison_by_Station"
    for regime in ["Perennial", "Intermittent"]:
        (base_dir / regime).mkdir(parents=True, exist_ok=True)

    for station, gs in pooling_summary.groupby("Station", sort=True):
        gs = gs.copy()
        regime = str(gs["Flow_Regime"].iloc[0])
        threshold_label = str(gs["Final_Threshold_Label"].iloc[0])
        threshold = float(gs["Final_Threshold_m3s"].iloc[0])
        ge = pooling_events.loc[pooling_events["Station"] == station].copy() if not pooling_events.empty else pd.DataFrame()

        ordered = gs.set_index("Pooling_Method").reindex(POOLING_METHOD_ORDER)
        labels = [POOLING_METHOD_LABELS[m] for m in POOLING_METHOD_ORDER]
        colors = [POOLING_COLORS[m] for m in POOLING_METHOD_ORDER]

        fig, axes = plt.subplots(1, 3, figsize=(14.2, 4.8))

        # A) Event frequency after pooling
        freq = pd.to_numeric(ordered["Complete_Events_per_Usable_Year"], errors="coerce").to_numpy(float)
        x = np.arange(len(POOLING_METHOD_ORDER))
        axes[0].bar(x, freq, color=colors, edgecolor="white", linewidth=0.4)
        axes[0].set_xticks(x)
        axes[0].set_xticklabels(labels, rotation=25, ha="right")
        axes[0].set_ylabel("Complete drought events per usable year")
        axes[0].set_title("Event frequency")

        # B-C) Complete-event distributions
        duration_data = []
        deficit_data = []
        n_complete = []
        for m in POOLING_METHOD_ORDER:
            if ge.empty:
                gd = pd.DataFrame()
            else:
                gd = ge.loc[(ge["Pooling_Method"] == m) & (ge["Censored"] == "No")].copy()
            dvals = pd.to_numeric(gd.get("Duration_days", pd.Series(dtype=float)), errors="coerce").dropna().to_numpy(float)
            vvals = pd.to_numeric(gd.get("Deficit_Volume_m3", pd.Series(dtype=float)), errors="coerce").dropna().to_numpy(float) / 1e6
            duration_data.append(dvals if len(dvals) else np.array([np.nan]))
            deficit_data.append(vvals if len(vvals) else np.array([np.nan]))
            n_complete.append(int(len(gd)))

        bp1 = axes[1].boxplot(duration_data, tick_labels=labels, patch_artist=True,
                              showfliers=False, widths=0.62,
                              medianprops={"color": "black", "linewidth": 1.2})
        bp2 = axes[2].boxplot(deficit_data, tick_labels=labels, patch_artist=True,
                              showfliers=False, widths=0.62,
                              medianprops={"color": "black", "linewidth": 1.2})
        for patch, c in zip(bp1["boxes"], colors):
            patch.set_facecolor(c); patch.set_alpha(0.70)
        for patch, c in zip(bp2["boxes"], colors):
            patch.set_facecolor(c); patch.set_alpha(0.70)

        axes[1].set_ylabel("Drought duration (days)")
        axes[1].set_title("Duration of complete events")
        axes[2].set_ylabel("Deficit volume (10$^6$ m$^3$)")
        axes[2].set_title("Deficit of complete events")

        for ax in axes[1:]:
            ax.tick_params(axis="x", labelrotation=25)
            for tick in ax.get_xticklabels():
                tick.set_ha("right")
        for j, n in enumerate(n_complete, start=1):
            axes[1].text(j, 0.985, f"n={n}", transform=axes[1].get_xaxis_transform(),
                         ha="center", va="top", fontsize=7)
            axes[2].text(j, 0.985, f"n={n}", transform=axes[2].get_xaxis_transform(),
                         ha="center", va="top", fontsize=7)

        for ax in axes:
            ax.grid(axis="y", linewidth=0.4, alpha=0.22)
            ax.set_axisbelow(True)
            ax.spines["top"].set_visible(False)
            ax.spines["right"].set_visible(False)

        fig.suptitle(
            f"{station} ({regime}) — fixed {threshold_label} = {threshold:.4g} m$^3$ s$^{{-1}}$",
            y=1.01, fontsize=12,
        )
        fig.text(
            0.5, 0.005,
            f"IT: tc={IT_CRITICAL_TIME_DAYS} d   |   "
            f"IC: tc={IC_CRITICAL_TIME_DAYS} d, pc={IC_CRITICAL_VOLUME_RATIO:.2f}   |   "
            f"MA: {MA_WINDOW_DAYS}-day centered moving average   |   SPA: sequent peak",
            ha="center", va="bottom", fontsize=8,
        )
        fig.tight_layout(rect=[0, 0.035, 1, 0.96])
        out = base_dir / regime / f"{station}_{threshold_label}_Pooling_Comparison.png"
        save_figure(fig, out)
        plt.close(fig)
        gc.collect()


def plot_pooling_hydrograph_station(
    daily: pd.DataFrame,
    station_events: pd.DataFrame,
    station: str,
    regime: str,
    threshold_label: str,
    threshold: float,
    fig_dir: Path,
) -> None:
    """Create one five-panel hydrograph per station for Raw, IT, IC, MA and SPA.

    The observed daily hydrograph is repeated on each panel so the different event
    definitions can be compared directly. The fixed Q90/Q50 threshold is shown as
    a dashed line. Coloured vertical spans show the duration assigned by each
    method, while the filled area below the threshold shows the physical streamflow
    deficit. Missing/unusable days remain NaN and therefore appear as breaks.
    """
    if daily.empty or station_events.empty or not np.isfinite(threshold) or threshold <= 0:
        return

    d = daily[["Time", "Flow"]].copy().reset_index(drop=True)
    d["Time"] = pd.to_datetime(d["Time"], errors="coerce")
    q = pd.to_numeric(d["Flow"], errors="coerce")
    valid = d["Time"].notna()
    d = d.loc[valid].reset_index(drop=True)
    q = pd.to_numeric(d["Flow"], errors="coerce")
    dates = pd.to_datetime(d["Time"])

    base_dir = fig_dir / "14_Final_Hydrographs_by_Station" / regime
    base_dir.mkdir(parents=True, exist_ok=True)

    fig, axes = plt.subplots(
        len(POOLING_METHOD_ORDER), 1, figsize=(15.5, 12.0), sharex=True, sharey=True
    )
    axes = np.asarray(axes).ravel()

    finite_q = q[np.isfinite(q)]
    if len(finite_q):
        q99 = float(finite_q.quantile(0.995))
        ymax = max(float(threshold) * 1.8, q99 * 1.08)
    else:
        ymax = float(threshold) * 2.0
    if not np.isfinite(ymax) or ymax <= 0:
        ymax = 1.0

    for ax, method in zip(axes, POOLING_METHOD_ORDER):
        color = POOLING_COLORS[method]
        ge = station_events.loc[station_events["Pooling_Method"] == method].copy()

        ax.plot(dates, q, color="#333333", linewidth=0.48, alpha=0.88, zorder=2)
        ax.axhline(
            threshold, color="#B2182B", linestyle="--", linewidth=1.15,
            label=f"{threshold_label} threshold", zorder=5
        )

        # Event duration = horizontal extent of the shaded span.
        # Deficit = filled vertical area between Q and Q_threshold where Q<Q_threshold.
        for _, ev in ge.iterrows():
            onset = pd.to_datetime(ev.get("Onset_Date"), errors="coerce")
            last = pd.to_datetime(ev.get("Last_Event_Day"), errors="coerce")
            if pd.isna(onset) or pd.isna(last):
                continue
            censored = str(ev.get("Censored", "No")) == "Yes"
            ax.axvspan(
                onset, last + pd.Timedelta(days=1),
                color=color, alpha=0.045 if censored else 0.10, zorder=0
            )

            mask = (dates >= onset) & (dates <= last) & q.notna() & (q < threshold)
            if mask.any():
                ax.fill_between(
                    dates, q, threshold, where=mask.to_numpy(dtype=bool),
                    interpolate=True, color=color, alpha=0.28, linewidth=0, zorder=1
                )

        complete = ge.loc[ge.get("Censored", pd.Series(index=ge.index, dtype=str)) == "No"].copy()
        med_d = pd.to_numeric(complete.get("Duration_days", pd.Series(dtype=float)), errors="coerce").median()
        med_v = pd.to_numeric(complete.get("Deficit_Volume_m3", pd.Series(dtype=float)), errors="coerce").median() / 1e6
        n_complete = len(complete)

        txt = f"{POOLING_METHOD_LABELS[method]}  |  n={n_complete}"
        if np.isfinite(med_d):
            txt += f"  |  median duration={med_d:.1f} d"
        if np.isfinite(med_v):
            txt += f"  |  median deficit={med_v:.3g} ×10⁶ m³"
        ax.text(
            0.006, 0.965, txt, transform=ax.transAxes, ha="left", va="top",
            fontsize=8.2, color="#111111",
            bbox=dict(boxstyle="round,pad=0.18", facecolor="white", edgecolor="none", alpha=0.82),
            zorder=8,
        )

        # Label at most the two longest complete events, so duration and deficit
        # are explicit without making multi-decade hydrographs unreadable.
        if not complete.empty:
            top = complete.copy()
            top["Duration_days"] = pd.to_numeric(top["Duration_days"], errors="coerce")
            top = top.sort_values("Duration_days", ascending=False).head(2)
            for _, ev in top.iterrows():
                onset = pd.to_datetime(ev.get("Onset_Date"), errors="coerce")
                last = pd.to_datetime(ev.get("Last_Event_Day"), errors="coerce")
                if pd.isna(onset) or pd.isna(last):
                    continue
                mid = onset + (last - onset) / 2
                dur = float(ev.get("Duration_days", np.nan))
                vol = float(ev.get("Deficit_Volume_m3", np.nan)) / 1e6
                label = f"D={dur:.0f} d"
                if np.isfinite(vol):
                    label += f"\nV={vol:.2g} Mm³"
                ax.annotate(
                    label, xy=(mid, threshold), xytext=(0, -24),
                    textcoords="offset points", ha="center", va="top", fontsize=6.4,
                    color=color, arrowprops=dict(arrowstyle="-", color=color, lw=0.6),
                    bbox=dict(boxstyle="round,pad=0.15", facecolor="white", edgecolor=color, alpha=0.75),
                    clip_on=True, zorder=9,
                )

        ax.set_ylim(bottom=0, top=ymax)
        ax.set_ylabel("Q (m$^3$ s$^{-1}$)")
        ax.grid(axis="y", linewidth=0.35, alpha=0.20)
        ax.set_axisbelow(True)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)

    axes[-1].set_xlabel("Date")
    fig.suptitle(
        f"{station} ({regime}) — hydrograph and drought events using fixed {threshold_label} = {threshold:.4g} m$^3$ s$^{{-1}}$",
        y=0.997, fontsize=12,
    )
    fig.text(
        0.5, 0.006,
        "Shaded horizontal span = event duration; filled area below threshold = physical streamflow deficit. "
        "Missing/unusable days remain gaps and are never pooled across.",
        ha="center", va="bottom", fontsize=8.2,
    )
    fig.tight_layout(rect=[0, 0.026, 1, 0.982])
    out = base_dir / f"{station}_{threshold_label}_Hydrograph_Raw_IT_IC_MA_SPA.png"
    save_figure(fig, out)
    plt.close(fig)
    gc.collect()


def _heatmap_metric(ax, matrix: pd.DataFrame, title: str, fmt: str = ".2g") -> None:
    """Draw one station x method heatmap used in the regime summary figures."""
    arr = matrix.to_numpy(dtype=float)
    finite = arr[np.isfinite(arr)]
    if finite.size == 0:
        ax.axis("off")
        ax.set_title(title + " (no data)")
        return
    im = ax.imshow(arr, aspect="auto", interpolation="nearest", cmap="viridis")
    ax.set_title(title)
    ax.set_xticks(np.arange(len(matrix.columns)))
    ax.set_xticklabels(matrix.columns, rotation=35, ha="right")
    ax.set_yticks(np.arange(len(matrix.index)))
    ax.set_yticklabels(matrix.index, fontsize=6.5)
    cbar = ax.figure.colorbar(im, ax=ax, pad=0.01, shrink=0.85)
    cbar.ax.tick_params(labelsize=7)


def plot_pooling_summary_all_stations(pooling_summary: pd.DataFrame, fig_dir: Path) -> None:
    """Create one all-stations summary figure for each flow regime.

    Each panel is a station x pooling-method heatmap. This keeps every station
    visible while showing the main frequency, duration and deficit information
    from the existing station-level comparison in one figure.
    """
    if pooling_summary.empty:
        return

    out_dir = fig_dir / "15_Final_Pooling_Summary_All_Stations"
    out_dir.mkdir(parents=True, exist_ok=True)

    metric_specs = [
        ("Complete_Events_per_Usable_Year", "Complete events per usable year", 1.0),
        ("Censored_Events", "Censored events", 1.0),
        ("Median_Duration_days", "Median duration (days)", 1.0),
        ("Mean_Duration_days", "Mean duration (days)", 1.0),
        ("Max_Duration_days", "Maximum duration (days)", 1.0),
        ("Median_Deficit_Volume_m3", "Median deficit volume (10$^6$ m$^3$)", 1e6),
        ("Mean_Deficit_Volume_m3", "Mean deficit volume (10$^6$ m$^3$)", 1e6),
        ("Max_Deficit_Volume_m3", "Maximum deficit volume (10$^6$ m$^3$)", 1e6),
    ]

    for regime in ["Perennial", "Intermittent"]:
        g = pooling_summary.loc[pooling_summary["Flow_Regime"] == regime].copy()
        if g.empty:
            continue

        stations = sorted(g["Station"].astype(str).unique().tolist())
        method_cols = [POOLING_METHOD_LABELS[m] for m in POOLING_METHOD_ORDER]
        fig_h = max(10.5, 0.19 * len(stations) + 7.0)
        fig, axes = plt.subplots(2, 4, figsize=(20.0, fig_h))
        axes = np.asarray(axes).ravel()

        for ax, (metric, title, scale) in zip(axes, metric_specs):
            temp = g[["Station", "Pooling_Method", metric]].copy()
            temp[metric] = pd.to_numeric(temp[metric], errors="coerce") / scale
            mat = temp.pivot(index="Station", columns="Pooling_Method", values=metric)
            mat = mat.reindex(index=stations, columns=POOLING_METHOD_ORDER)
            mat.columns = method_cols
            _heatmap_metric(ax, mat, title)
            ax.set_ylabel("Station")

        thr = "Q90" if regime == "Perennial" else "Q50"
        fig.suptitle(
            f"{regime} stations — fixed {thr}: Raw, IT, IC, MA and SPA comparison",
            y=0.995, fontsize=13,
        )
        fig.text(
            0.5, 0.008,
            "Rows are individual stations; columns are pooling approaches. Duration and deficit statistics use complete (uncensored) events.",
            ha="center", va="bottom", fontsize=8.5,
        )
        fig.tight_layout(rect=[0, 0.022, 1, 0.975])
        save_figure(fig, out_dir / f"{regime}_All_Stations_Pooling_Summary.png")
        plt.close(fig)
        gc.collect()


def load_metadata(metadata_path: Path | None, input_dir: Path) -> pd.DataFrame | None:
    candidates: list[Path] = []
    if metadata_path is not None:
        candidates.append(metadata_path)
    candidates.extend([
        input_dir.parent / "Catchments_info.csv",
        input_dir.parent.parent / "Catchments_info.csv",
        Path("/mnt/data/Catchments_info.csv"),
    ])

    for p in candidates:
        if p.exists():
            try:
                m = pd.read_csv(p, encoding="utf-8-sig")
                if "Station_Code" in m.columns:
                    m["Station_Code"] = m["Station_Code"].astype(str).str.strip()
                    return m
            except Exception:
                continue
    return None


def metadata_for_station(meta: pd.DataFrame | None, station: str) -> dict:
    if meta is None:
        return {}
    hit = meta.loc[meta["Station_Code"] == station]
    if hit.empty:
        return {}
    r = hit.iloc[0]
    return {
        "Original_Station_ID": r.get("Station_ID", np.nan),
        "Catchment": r.get("Catchment", np.nan),
        "Main_River": r.get("Main river", np.nan),
        "Compartment": r.get("Compartment", np.nan),
    }


# =============================================================================
# DAILY QC
# =============================================================================
def apply_isolated_zero_qc(df: pd.DataFrame) -> tuple[pd.DataFrame, list[dict]]:
    """Apply the agreed isolated-zero rule on a complete daily calendar."""
    out = df.copy()
    flow = pd.to_numeric(out["Flow"], errors="coerce").copy()
    original = flow.copy()
    zero = original.eq(0).to_numpy(dtype=bool)
    log: list[dict] = []

    if not zero.any():
        out["Flow"] = flow
        return out, log

    starts = np.where(zero & np.r_[True, ~zero[:-1]])[0]
    ends = np.where(zero & np.r_[~zero[1:], True])[0]

    for s, e in zip(starts, ends):
        run_len = int(e - s + 1)
        if run_len >= 2:
            continue

        i = int(s)
        prev_val = original.iloc[i - 1] if i > 0 else np.nan
        next_val = original.iloc[i + 1] if i < len(original) - 1 else np.nan
        date = out.iloc[i]["Time"]

        if pd.notna(prev_val) and pd.notna(next_val):
            replacement = (float(prev_val) + float(next_val)) / 2.0
            flow.iloc[i] = replacement
            action = "Replaced_by_neighbor_mean"
        else:
            replacement = np.nan
            flow.iloc[i] = np.nan
            action = "Converted_to_NaN"

        log.append({
            "Date": pd.Timestamp(date).date().isoformat(),
            "Original_Flow": 0.0,
            "Previous_Day_Flow": prev_val,
            "Next_Day_Flow": next_val,
            "Action": action,
            "Replacement_Flow": replacement,
        })

    out["Flow"] = flow
    return out, log


def read_station_csv(path: Path) -> tuple[pd.DataFrame, dict, list[dict]]:
    """Read, standardise, de-duplicate, calendar-complete, and apply zero QC."""
    raw = pd.read_csv(path, encoding="utf-8-sig")
    col_lookup = {str(c).strip().lower(): c for c in raw.columns}
    if "time" not in col_lookup or "flow" not in col_lookup:
        raise ValueError(f"{path.name}: expected columns Time and Flow; found {list(raw.columns)}")

    df = raw[[col_lookup["time"], col_lookup["flow"]]].copy()
    df.columns = ["Time", "Flow"]

    n_rows_raw = len(df)
    df["Time"] = pd.to_datetime(df["Time"], errors="coerce")
    n_bad_dates = int(df["Time"].isna().sum())
    df["Flow"] = pd.to_numeric(df["Flow"], errors="coerce")

    n_sentinel = int(df["Flow"].isin([-999, -9999]).sum())
    df.loc[df["Flow"].isin([-999, -9999]), "Flow"] = np.nan
    n_negative = int((df["Flow"] < 0).sum())
    df.loc[df["Flow"] < 0, "Flow"] = np.nan

    df = df.dropna(subset=["Time"]).sort_values("Time")

    duplicate_rows = int(df["Time"].duplicated(keep=False).sum())
    duplicate_dates = int(df.loc[df["Time"].duplicated(keep=False), "Time"].nunique())
    duplicate_conflicts = 0

    if duplicate_dates > 0:
        collapsed = []
        for date, g in df.groupby("Time", sort=True):
            vals = pd.to_numeric(g["Flow"], errors="coerce").dropna().unique()
            if len(vals) == 0:
                value = np.nan
            elif len(vals) == 1:
                value = float(vals[0])
            else:
                value = np.nan
                duplicate_conflicts += 1
            collapsed.append((date, value))
        df = pd.DataFrame(collapsed, columns=["Time", "Flow"])

    if df.empty:
        qc = {
            "Raw_Rows": n_rows_raw,
            "Unparseable_Dates": n_bad_dates,
            "Sentinel_Values": n_sentinel,
            "Negative_Flows": n_negative,
            "Duplicate_Rows": duplicate_rows,
            "Duplicate_Dates": duplicate_dates,
            "Duplicate_Conflict_Dates": duplicate_conflicts,
            "Inserted_Calendar_Days": 0,
        }
        return df, qc, []

    original_unique_days = int(df["Time"].nunique())
    full_index = pd.date_range(df["Time"].min(), df["Time"].max(), freq="D")
    df = df.set_index("Time").reindex(full_index)
    df.index.name = "Time"
    df = df.reset_index()
    inserted_calendar_days = int(len(df) - original_unique_days)

    df, zero_log = apply_isolated_zero_qc(df)

    qc = {
        "Raw_Rows": n_rows_raw,
        "Unparseable_Dates": n_bad_dates,
        "Sentinel_Values": n_sentinel,
        "Negative_Flows": n_negative,
        "Duplicate_Rows": duplicate_rows,
        "Duplicate_Dates": duplicate_dates,
        "Duplicate_Conflict_Dates": duplicate_conflicts,
        "Inserted_Calendar_Days": inserted_calendar_days,
        "Isolated_Zeros_Averaged": sum(x["Action"] == "Replaced_by_neighbor_mean" for x in zero_log),
        "Isolated_Zeros_Converted_to_NaN": sum(x["Action"] == "Converted_to_NaN" for x in zero_log),
    }
    return df, qc, zero_log


# =============================================================================
# STATION ANALYSIS
# =============================================================================
def analyse_station(path: Path, meta: pd.DataFrame | None) -> dict | None:
    station = path.stem.strip()
    archive, qc, zero_qc_log = read_station_csv(path)

    if archive.empty or not archive["Flow"].notna().any():
        return None

    first_valid = archive.loc[archive["Flow"].notna(), "Time"].min()
    last_valid = archive.loc[archive["Flow"].notna(), "Time"].max()
    record = archive.loc[(archive["Time"] >= first_valid) &
                         (archive["Time"] <= last_valid)].copy().reset_index(drop=True)

    valid = record["Flow"].notna()
    missing = ~valid
    zero = valid & np.isclose(record["Flow"], 0.0, atol=0.0)

    valid_days_all = int(valid.sum())
    record_days = int(len(record))
    internal_missing_days = int(missing.sum())
    internal_missing_pct = 100.0 * internal_missing_days / record_days if record_days else np.nan
    effective_valid_years = valid_days_all / 365.0

    missing_segments = run_segments(missing)
    valid_segments = run_segments(valid)
    zero_segments_all = run_segments(zero)

    longest_gap = max((e - s + 1 for s, e in missing_segments), default=0)
    longest_valid_block = max((e - s + 1 for s, e in valid_segments), default=0)
    longest_zero_all = max((e - s + 1 for s, e in zero_segments_all), default=0)

    # -------------------------------------------------------------------------
    # Annual completeness: <=15 missing days, including correct leap-year rule
    # -------------------------------------------------------------------------
    annual_rows = []
    usable_years: list[int] = []

    for year in range(int(first_valid.year), int(last_valid.year) + 1):
        calendar_days = 366 if calendar.isleap(year) else 365
        y0 = pd.Timestamp(year=year, month=1, day=1)
        y1 = pd.Timestamp(year=year, month=12, day=31)
        g = archive.loc[(archive["Time"] >= y0) & (archive["Time"] <= y1)].copy()

        valid_days = int(g["Flow"].notna().sum())
        missing_days = int(calendar_days - valid_days)
        zero_days = int((g["Flow"].notna() & np.isclose(g["Flow"], 0.0, atol=0.0)).sum())
        usable = missing_days <= MAX_MISSING_DAYS_PER_USABLE_YEAR

        if usable:
            usable_years.append(year)

        annual_rows.append({
            "Station": station,
            "Year": year,
            "Calendar_Days": calendar_days,
            "Valid_Days": valid_days,
            "Missing_Days": missing_days,
            "Coverage_pct": 100.0 * valid_days / calendar_days,
            "Zero_Flow_Days": zero_days,
            "Zero_Flow_pct_of_Valid_Days": 100.0 * zero_days / valid_days if valid_days else np.nan,
            "Usable_Year": "Yes" if usable else "No",
            "Annual_Completeness_Rule": "Missing days <= 15",
        })

    annual_df = pd.DataFrame(annual_rows)
    n_usable_years = len(usable_years)
    retained = n_usable_years >= MIN_USABLE_YEARS

    # Data used for classification and FDC thresholds
    usable_mask = archive["Time"].dt.year.isin(usable_years)
    usable_daily = archive.loc[usable_mask].copy()
    usable_valid = usable_daily["Flow"].notna()
    usable_flow = usable_daily.loc[usable_valid, "Flow"]

    usable_valid_days = int(usable_valid.sum())
    usable_zero_days = int((usable_valid & np.isclose(usable_daily["Flow"], 0.0, atol=0.0)).sum())
    zero_pct_usable = 100.0 * usable_zero_days / usable_valid_days if usable_valid_days else np.nan

    # Additional annual diagnostic, not the classification variable
    mean_zero_days_per_usable_year = (
        usable_zero_days / n_usable_years if n_usable_years else np.nan
    )

    if usable_valid_days == 0:
        regime = "Not assessed"
    elif zero_pct_usable < REGIME_ZERO_PCT_CUTOFF:
        regime = "Perennial"
    else:
        regime = "Intermittent"

    # Zero spells using usable years only; unusable years become gaps
    usable_series_for_spells = archive.copy()
    usable_series_for_spells.loc[
        ~usable_series_for_spells["Time"].dt.year.isin(usable_years), "Flow"
    ] = np.nan
    uzero = usable_series_for_spells["Flow"].notna() & np.isclose(
        usable_series_for_spells["Flow"], 0.0, atol=0.0
    )
    zero_segments_usable = run_segments(uzero)
    longest_zero_usable = max((e - s + 1 for s, e in zero_segments_usable), default=0)

    summary = {
        "Station": station,
        **metadata_for_station(meta, station),
        "First_Valid_Date": first_valid.date().isoformat(),
        "Last_Valid_Date": last_valid.date().isoformat(),
        "Historical_Period": f"{first_valid.year}-{last_valid.year}",
        "Calendar_Years_Spanned": int(last_valid.year - first_valid.year + 1),
        "Valid_Daily_Observations_AllYears": valid_days_all,
        "Effective_Valid_Years_AllYears_validDaysDiv365": effective_valid_years,
        "Usable_Years": n_usable_years,
        "NonUsable_Years": int(len(annual_df) - n_usable_years),
        "Meets_Minimum_Usable_Years": "Yes" if retained else "No",
        "Main_Analysis_Status": "Retained" if retained else "Excluded",
        "Usable_Valid_Days": usable_valid_days,
        "Internal_Missing_Days": internal_missing_days,
        "Internal_Missing_pct": internal_missing_pct,
        "Number_of_Internal_Missing_Periods": len(missing_segments),
        "Longest_Internal_Gap_days": longest_gap,
        "Number_of_Continuous_Valid_Blocks": len(valid_segments),
        "Longest_Continuous_Valid_Block_days": longest_valid_block,
        "Zero_Flow_Days_AllYears": int(zero.sum()),
        "Zero_Flow_pct_AllValidDays": 100.0 * int(zero.sum()) / valid_days_all if valid_days_all else np.nan,
        "Longest_Zero_Flow_Spell_days_AllYears": longest_zero_all,
        "Zero_Flow_Days_UsableYears": usable_zero_days,
        "Zero_Flow_pct_UsableYears": zero_pct_usable,
        "Mean_Zero_Flow_Days_per_Usable_Year": mean_zero_days_per_usable_year,
        "Number_of_Zero_Flow_Spells_UsableYears": len(zero_segments_usable),
        "Longest_Zero_Flow_Spell_days_UsableYears": longest_zero_usable,
        "Flow_Regime": regime,
        "Regime_ZeroFlow_pct_Cutoff": REGIME_ZERO_PCT_CUTOFF,
        "Annual_Completeness_Rule": "<=15 missing days per calendar year",
        "Station_Inclusion_Rule": f">={MIN_USABLE_YEARS} usable years",
        **qc,
    }

    # Missing-gap audit table
    missing_rows = []
    for j, (s, e) in enumerate(missing_segments, start=1):
        missing_rows.append({
            "Station": station,
            "Gap_ID": j,
            "Start_Date": record.loc[s, "Time"].date().isoformat(),
            "End_Date": record.loc[e, "Time"].date().isoformat(),
            "Duration_days": int(e - s + 1),
        })

    # Continuous-valid-block audit table
    block_rows = []
    for j, (s, e) in enumerate(valid_segments, start=1):
        block_rows.append({
            "Station": station,
            "Block_ID": j,
            "Start_Date": record.loc[s, "Time"].date().isoformat(),
            "End_Date": record.loc[e, "Time"].date().isoformat(),
            "Duration_days": int(e - s + 1),
            "Equivalent_years": (e - s + 1) / 365.0,
        })

    # Zero-spell audit table within usable years only
    zero_rows = []
    for j, (s, e) in enumerate(zero_segments_usable, start=1):
        zero_rows.append({
            "Station": station,
            "Zero_Spell_ID": j,
            "Start_Date": usable_series_for_spells.loc[s, "Time"].date().isoformat(),
            "End_Date": usable_series_for_spells.loc[e, "Time"].date().isoformat(),
            "Duration_days": int(e - s + 1),
        })

    # QC log for isolated zero changes
    zero_qc_rows = []
    for row in zero_qc_log:
        zero_qc_rows.append({"Station": station, **row})

    # Threshold tables only for retained stations
    perennial_thresholds = None
    intermittent_thresholds = None
    event_sensitivity_rows: list[dict] = []
    event_catalog_rows: list[dict] = []
    pooling_summary_rows: list[dict] = []
    pooling_event_rows: list[dict] = []

    # Daily analysis series used for all threshold/event comparisons.
    # Unusable years are masked to NaN; existing gaps remain NaN.
    event_daily = archive[["Time", "Flow"]].copy()
    event_daily.loc[~event_daily["Time"].dt.year.isin(usable_years), "Flow"] = np.nan

    if retained and usable_valid_days > 0:
        if regime == "Perennial":
            q95 = fdc_q(usable_flow, 95)
            q90 = fdc_q(usable_flow, 90)
            q85 = fdc_q(usable_flow, 85)
            perennial_thresholds = {
                "Station": station,
                "Usable_Years": n_usable_years,
                "Usable_Valid_Days": usable_valid_days,
                "Zero_Flow_pct_UsableYears": zero_pct_usable,
                "Q95_m3s": q95,
                "Q90_m3s": q90,
                "Q85_m3s": q85,
                "Selected_Threshold": "Q90" if pd.notna(q90) and q90 > 0 else "None",
                "Selected_Threshold_m3s": q90 if pd.notna(q90) and q90 > 0 else np.nan,
                "Days_Below_Q95": int((usable_flow < q95).sum()),
                "Days_Below_Q90": int((usable_flow < q90).sum()),
                "Days_Below_Q85": int((usable_flow < q85).sum()),
                "Pct_Below_Q95": 100.0 * (usable_flow < q95).sum() / usable_valid_days,
                "Pct_Below_Q90": 100.0 * (usable_flow < q90).sum() / usable_valid_days,
                "Pct_Below_Q85": 100.0 * (usable_flow < q85).sum() / usable_valid_days,
            }
        elif regime == "Intermittent":
            # Full intermittent-stream FDC sensitivity set.
            q70 = fdc_q(usable_flow, 70)
            q60 = fdc_q(usable_flow, 60)
            q50 = fdc_q(usable_flow, 50)
            q40 = fdc_q(usable_flow, 40)
            q30 = fdc_q(usable_flow, 30)
            q20 = fdc_q(usable_flow, 20)

            # Final intermittent drought analysis uses the fixed Q50 threshold.
            # Other percentiles remain available only as sensitivity diagnostics.
            if pd.notna(q50) and q50 > 0:
                selected_name, selected_q = "Q50", q50
            else:
                selected_name, selected_q = "None", np.nan

            intermittent_thresholds = {
                "Station": station,
                "Usable_Years": n_usable_years,
                "Usable_Valid_Days": usable_valid_days,
                "Zero_Flow_Days_UsableYears": usable_zero_days,
                "Zero_Flow_pct_UsableYears": zero_pct_usable,
                "Q70_m3s": q70,
                "Q60_m3s": q60,
                "Q50_m3s": q50,
                "Q40_m3s": q40,
                "Q30_m3s": q30,
                "Q20_m3s": q20,
                "Selected_Threshold": selected_name,
                "Selected_Threshold_m3s": selected_q,
            }

        # Event-based sensitivity for EVERY candidate threshold.
        if regime == "Perennial" and perennial_thresholds is not None:
            candidate_thresholds = [
                ("Q95", 95.0, perennial_thresholds["Q95_m3s"]),
                ("Q90", 90.0, perennial_thresholds["Q90_m3s"]),
                ("Q85", 85.0, perennial_thresholds["Q85_m3s"]),
            ]
        elif regime == "Intermittent" and intermittent_thresholds is not None:
            candidate_thresholds = [
                ("Q70", 70.0, intermittent_thresholds["Q70_m3s"]),
                ("Q60", 60.0, intermittent_thresholds["Q60_m3s"]),
                ("Q50", 50.0, intermittent_thresholds["Q50_m3s"]),
                ("Q40", 40.0, intermittent_thresholds["Q40_m3s"]),
                ("Q30", 30.0, intermittent_thresholds["Q30_m3s"]),
                ("Q20", 20.0, intermittent_thresholds["Q20_m3s"]),
            ]
        else:
            candidate_thresholds = []

        for label, exc, qthr in candidate_thresholds:
            sens_row, ev_rows = threshold_event_summary(
                event_daily, qthr, station, regime, label, exc, n_usable_years
            )
            event_sensitivity_rows.append(sens_row)
            event_catalog_rows.extend(ev_rows)

        # ---------------------------------------------------------------------
        # FINAL drought analysis requested for the study:
        #   Perennial    -> fixed Q90
        #   Intermittent -> fixed Q50
        # Compare Raw, IT, IC, MA and SPA on the SAME station threshold.
        # ---------------------------------------------------------------------
        if regime == "Perennial" and perennial_thresholds is not None:
            final_label = "Q90"
            final_q = perennial_thresholds.get("Q90_m3s", np.nan)
        elif regime == "Intermittent" and intermittent_thresholds is not None:
            final_label = "Q50"
            final_q = intermittent_thresholds.get("Q50_m3s", np.nan)
        else:
            final_label = "None"
            final_q = np.nan

        summary["Final_Drought_Threshold_Label"] = final_label
        summary["Final_Drought_Threshold_m3s"] = final_q

        if np.isfinite(final_q) and final_q > 0:
            ps, pe = final_pooling_analysis(
                event_daily, final_q, station, regime, final_label, n_usable_years
            )
            pooling_summary_rows.extend(ps)
            pooling_event_rows.extend(pe)

    return {
        "summary": summary,
        "annual": annual_df,
        "missing": missing_rows,
        "blocks": block_rows,
        "zero_spells": zero_rows,
        "zero_qc": zero_qc_rows,
        "perennial": perennial_thresholds,
        "intermittent": intermittent_thresholds,
        "event_sensitivity": event_sensitivity_rows,
        "event_catalog": event_catalog_rows,
        "pooling_summary": pooling_summary_rows,
        "pooling_events": pooling_event_rows,
        # Returned only so main() can make the station hydrograph immediately;
        # it is not accumulated across stations or written to Excel.
        "event_daily": event_daily if retained else pd.DataFrame(),
    }


# =============================================================================
# FIGURES
# =============================================================================
def barh_plot(df: pd.DataFrame, value_col: str, xlabel: str, path: Path,
              cutoff: float | None = None, cutoff_text: str | None = None,
              title: str | None = None, colors=None) -> None:
    if df.empty:
        return

    d = df.sort_values(value_col, ascending=True).reset_index(drop=True)
    h = max(6.0, 0.22 * len(d) + 1.5)
    fig, ax = plt.subplots(figsize=(9.5, h))

    if colors is None:
        bar_colors = COLOR_PERENNIAL
    elif callable(colors):
        bar_colors = colors(d)
    else:
        bar_colors = colors

    ax.barh(d["Station"], d[value_col], color=bar_colors, edgecolor="white", linewidth=0.25)

    if cutoff is not None:
        ax.axvline(cutoff, linestyle="--", linewidth=1.1, color="#222222", zorder=4)
        if cutoff_text:
            ymax = len(d) - 0.5
            ax.text(cutoff, ymax, "  " + cutoff_text, ha="left", va="top", fontsize=8,
                    color="#222222")

    if title:
        ax.set_title(title)
    ax.set_xlabel(xlabel)
    ax.set_ylabel("Station")
    ax.grid(axis="x", linewidth=0.4, alpha=0.22)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    save_figure(fig, path)
    plt.close(fig)
    gc.collect()
def _plot_metric_with_iqr(ax, summary: pd.DataFrame, labels: list[str],
                          metric_base: str, ylabel: str, color: str) -> None:
    d = summary.set_index("Threshold_Label").reindex(labels)
    x = np.arange(len(labels))
    med = pd.to_numeric(d[f"{metric_base}_Median"], errors="coerce").to_numpy(float)
    q25 = pd.to_numeric(d[f"{metric_base}_Q25"], errors="coerce").to_numpy(float)
    q75 = pd.to_numeric(d[f"{metric_base}_Q75"], errors="coerce").to_numpy(float)
    ax.plot(x, med, marker="o", linewidth=1.6, markersize=4.5, color=color)
    ax.fill_between(x, q25, q75, alpha=0.18, color=color, linewidth=0)
    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    ax.set_ylabel(ylabel)
    ax.grid(axis="y", linewidth=0.4, alpha=0.22)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)


def plot_event_sensitivity(regime_summary: pd.DataFrame, regime: str,
                           labels: list[str], path: Path) -> None:
    """Event-sensitivity diagnostic using station medians and IQRs."""
    d = regime_summary.loc[regime_summary["Flow_Regime"] == regime].copy()
    if d.empty:
        return

    if regime == "Perennial":
        fig, axes = plt.subplots(2, 2, figsize=(10.5, 7.4))
        specs = [
            ("Drought_Days_per_Usable_Year", "Drought days per usable year", COLOR_Q90),
            ("Complete_Events_per_Year", "Complete events per usable year", COLOR_Q60),
            ("Median_Duration_days_Complete", "Median event duration (days)", COLOR_Q40),
            ("Median_MinFlow_Deficit_Ratio_Complete", "Median min-flow deficit ratio", COLOR_Q70),
        ]
    else:
        fig, axes = plt.subplots(2, 3, figsize=(13.0, 7.6))
        specs = [
            ("Drought_Days_per_Usable_Year", "Drought days per usable year", COLOR_Q90),
            ("Complete_Events_per_Year", "Complete events per usable year", COLOR_Q60),
            ("Median_Duration_days_Complete", "Median event duration (days)", COLOR_Q40),
            ("Median_MinFlow_Deficit_Ratio_Complete", "Median min-flow deficit ratio", COLOR_Q70),
            ("Zero_Flow_Share_of_Drought_Days_pct", "Zero-flow share of drought days (%)", COLOR_Q85),
            ("Censored_Event_Share_pct", "Censored event share (%)", COLOR_Q30),
        ]

    axes = np.asarray(axes).ravel()
    for ax, (metric, ylabel, color) in zip(axes, specs):
        _plot_metric_with_iqr(ax, d, labels, metric, ylabel, color)

    fig.suptitle(f"{regime} threshold sensitivity of drought-event characteristics", y=0.995)
    fig.tight_layout(rect=[0, 0, 1, 0.975])
    save_figure(fig, path)
    plt.close(fig)
    gc.collect()

def plot_timing_heatmaps(monthly_df: pd.DataFrame, regime: str, labels: list[str],
                         path: Path) -> None:
    """Two-panel onset/recovery month heatmap for complete events."""
    d = monthly_df.loc[monthly_df["Flow_Regime"] == regime].copy()
    if d.empty:
        return

    month_names = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
                   "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    fig, axes = plt.subplots(
        1, 2,
        figsize=(13.0, max(4.2, 0.55 * len(labels) + 2.2)),
        constrained_layout=True,
    )
    im = None
    for ax, timing in zip(axes, ["Onset", "Recovery"]):
        t = d.loc[d["Timing"] == timing].pivot_table(
            index="Threshold_Label", columns="Month", values="Event_Percentage",
            aggfunc="first"
        ).reindex(labels).reindex(columns=range(1, 13))
        arr = t.to_numpy(float)
        im = ax.imshow(arr, aspect="auto", interpolation="nearest", cmap="YlOrRd", vmin=0)
        ax.set_xticks(np.arange(12))
        ax.set_xticklabels(month_names, rotation=45, ha="right")
        ax.set_yticks(np.arange(len(labels)))
        ax.set_yticklabels(labels)
        ax.set_xlabel("Month")
        ax.set_ylabel("Threshold")
        ax.set_title(f"{timing} month")
    if im is not None:
        cbar = fig.colorbar(im, ax=axes, pad=0.02, shrink=0.9)
        cbar.set_label("Complete events (%)")
    fig.suptitle(f"{regime} event timing sensitivity")
    save_figure(fig, path)
    plt.close(fig)
    gc.collect()


def make_plots(all_df: pd.DataFrame, retained: pd.DataFrame,
               annual_ret: pd.DataFrame, per_df: pd.DataFrame,
               int_df: pd.DataFrame, regime_sensitivity: pd.DataFrame,
               monthly_timing: pd.DataFrame, fig_dir: Path) -> None:
    """Create publication-quality color figures.

    The previous zoomed figure around the regime boundary is intentionally
    omitted. Regime information is shown only in the main zero-flow figure.
    """
    fig_dir.mkdir(parents=True, exist_ok=True)

    # 1. Usable years by station: retained vs excluded
    d = all_df.sort_values("Usable_Years", ascending=True).reset_index(drop=True)
    h = max(8.0, 0.20 * len(d) + 1.8)
    fig, ax = plt.subplots(figsize=(9.5, h))
    colors = [COLOR_RETAINED if x >= MIN_USABLE_YEARS else COLOR_EXCLUDED
              for x in d["Usable_Years"]]
    ax.barh(d["Station"], d["Usable_Years"], color=colors, edgecolor="white", linewidth=0.25)
    ax.axvline(MIN_USABLE_YEARS, linestyle="--", linewidth=1.1, color="#222222")
    ax.text(MIN_USABLE_YEARS, len(d) - 0.5, "  minimum record criterion",
            ha="left", va="top", fontsize=8, color="#222222")
    ax.set_xlabel("Number of usable calendar years")
    ax.set_ylabel("Station")
    ax.grid(axis="x", linewidth=0.4, alpha=0.22)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    save_figure(fig, fig_dir / "01_Usable_Years_by_Station.png")
    plt.close(fig)
    gc.collect()
    # 2. Historical period: color by flow regime
    if not retained.empty:
        d = retained.copy()
        d["StartYear"] = pd.to_datetime(d["First_Valid_Date"]).dt.year
        d["EndYear"] = pd.to_datetime(d["Last_Valid_Date"]).dt.year
        d = d.sort_values(["StartYear", "EndYear"]).reset_index(drop=True)
        h = max(6.5, 0.24 * len(d) + 1.5)
        fig, ax = plt.subplots(figsize=(10.5, h))
        y = np.arange(len(d))
        line_colors = [COLOR_PERENNIAL if r == "Perennial" else COLOR_INTERMITTENT
                       for r in d["Flow_Regime"]]
        for yi, start_y, end_y, c in zip(y, d["StartYear"], d["EndYear"], line_colors):
            ax.hlines(yi, start_y, end_y, color=c, linewidth=2.5, alpha=0.9)
            ax.scatter(start_y, yi, s=14, color=c, zorder=3)
            ax.scatter(end_y, yi, s=18, facecolors="white", edgecolors=c,
                       linewidth=0.9, zorder=3)
        ax.set_yticks(y)
        ax.set_yticklabels(d["Station"])
        ax.set_xlabel("Calendar year")
        ax.set_ylabel("Station")
        ax.grid(axis="x", linewidth=0.4, alpha=0.22)
        ax.set_axisbelow(True)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        from matplotlib.patches import Patch
        ax.legend(handles=[
            Patch(facecolor=COLOR_PERENNIAL, edgecolor="none", label="Perennial"),
            Patch(facecolor=COLOR_INTERMITTENT, edgecolor="none", label="Intermittent"),
        ], frameon=False, loc="best", ncol=2)
        fig.tight_layout()
        save_figure(fig, fig_dir / "02_Historical_Periods_Retained.png")
        plt.close(fig)
        gc.collect()
    # 3. Internal missing data: continuous color gradient
    if not retained.empty:
        dmiss = retained.sort_values("Internal_Missing_pct", ascending=True).reset_index(drop=True)
        vals = dmiss["Internal_Missing_pct"].to_numpy(float)
        if len(vals):
            vmin = np.nanmin(vals)
            vmax = np.nanmax(vals)
            if not np.isfinite(vmin):
                vmin = 0.0
            if not np.isfinite(vmax) or vmax <= vmin:
                vmax = vmin + 1.0
            norm = plt.Normalize(vmin=vmin, vmax=vmax)
            cmap = plt.get_cmap("viridis")
            bar_colors = cmap(norm(vals))
            h = max(6.0, 0.22 * len(dmiss) + 1.5)
            fig, ax = plt.subplots(figsize=(9.5, h))
            ax.barh(dmiss["Station"], vals, color=bar_colors, edgecolor="white", linewidth=0.25)
            ax.set_xlabel("Internal missing data (%)")
            ax.set_ylabel("Station")
            ax.grid(axis="x", linewidth=0.4, alpha=0.22)
            ax.set_axisbelow(True)
            ax.spines["top"].set_visible(False)
            ax.spines["right"].set_visible(False)
            sm = plt.cm.ScalarMappable(norm=norm, cmap=cmap)
            sm.set_array([])
            cbar = fig.colorbar(sm, ax=ax, pad=0.01)
            cbar.set_label("Missing data (%)")
            fig.tight_layout()
            save_figure(fig, fig_dir / "03_Internal_Missing_Data_pct.png")
            plt.close(fig)
            gc.collect()
    # 4. Zero-flow percentage and regime boundary, colored by classification
    if not retained.empty:
        def regime_colors(dframe: pd.DataFrame):
            return [COLOR_PERENNIAL if r == "Perennial" else COLOR_INTERMITTENT
                    for r in dframe["Flow_Regime"]]

        # Draw directly here so the classification colors are explicit.
        # No legend is used in Figure 04.
        dz = retained.sort_values("Zero_Flow_pct_UsableYears", ascending=True).reset_index(drop=True)
        hz = max(6.0, 0.22 * len(dz) + 1.5)
        fig, ax = plt.subplots(figsize=(9.5, hz))
        zcolors = regime_colors(dz)
        ax.barh(dz["Station"], dz["Zero_Flow_pct_UsableYears"],
                color=zcolors, edgecolor="white", linewidth=0.25)
        ax.axvline(REGIME_ZERO_PCT_CUTOFF, linestyle="--", linewidth=1.1,
                   color="#222222", zorder=4)
        ax.text(REGIME_ZERO_PCT_CUTOFF, len(dz) - 0.5,
                f"  {REGIME_ZERO_PCT_CUTOFF:.3f}%",
                ha="left", va="top", fontsize=8, color="#222222")
        ax.set_xlabel("Zero-flow days (% of valid observations in usable years)")
        ax.set_ylabel("Station")
        ax.grid(axis="x", linewidth=0.4, alpha=0.22)
        ax.set_axisbelow(True)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        fig.tight_layout()
        save_figure(fig, fig_dir / "04_ZeroFlow_Percentage_and_Regime_Threshold.png")
        plt.close(fig)
        gc.collect()
    # 5. Perennial threshold sensitivity: colored Q95/Q90/Q85 with legend
    if not per_df.empty:
        p = per_df.sort_values("Q90_m3s").reset_index(drop=True)
        x = np.arange(len(p))
        fig, ax = plt.subplots(figsize=(max(10.5, 0.36 * len(p) + 3.5), 6.0))

        ax.plot(x, p["Q95_m3s"], marker="o", markersize=3.0, linewidth=1.1,
                linestyle="--", color=COLOR_Q95, label="Q95")
        ax.plot(x, p["Q90_m3s"], marker="s", markersize=3.2, linewidth=1.4,
                linestyle="-", color=COLOR_Q90, label="Q90")
        ax.plot(x, p["Q85_m3s"], marker="^", markersize=3.2, linewidth=1.1,
                linestyle=":", color=COLOR_Q85, label="Q85")

        ax.set_xticks(x)
        ax.set_xticklabels(p["Station"], rotation=90)
        ax.set_ylabel("Discharge (m$^3$ s$^{-1}$)")
        ax.set_xlabel("Perennial station")
        ax.grid(axis="y", linewidth=0.4, alpha=0.22)
        ax.set_axisbelow(True)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        ax.legend(frameon=False, loc="upper left", ncol=3, fontsize=11, markerscale=1.25)

        fig.tight_layout()
        save_figure(fig, fig_dir / "05_Perennial_FDC_Thresholds.png")
        plt.close(fig)
        gc.collect()
    # 6. Intermittent threshold sensitivity: Q70, Q60, Q50, Q40, Q30, Q20
    #    This mirrors the perennial threshold-sensitivity figure.
    if not int_df.empty:
        im = int_df.sort_values("Q50_m3s").reset_index(drop=True)
        x = np.arange(len(im))
        fig, ax = plt.subplots(figsize=(max(10.5, 0.40 * len(im) + 3.5), 6.2))

        series = [
            ("Q70_m3s", "Q70", COLOR_Q70, "o", "--"),
            ("Q60_m3s", "Q60", COLOR_Q60, "s", "-"),
            ("Q50_m3s", "Q50", COLOR_Q50, "^", "-."),
            ("Q40_m3s", "Q40", COLOR_Q40, "D", ":"),
            ("Q30_m3s", "Q30", COLOR_Q30, "v", "--"),
            ("Q20_m3s", "Q20", COLOR_Q20, "P", "-"),
        ]
        for col, label, color, marker, ls in series:
            ax.plot(
                x, im[col], marker=marker, markersize=3.4, linewidth=1.25,
                linestyle=ls, color=color, label=label
            )

        ax.set_xticks(x)
        ax.set_xticklabels(im["Station"], rotation=90)
        ax.set_ylabel("Discharge (m$^3$ s$^{-1}$)")
        ax.set_xlabel("Intermittent station")
        ax.grid(axis="y", linewidth=0.4, alpha=0.22)
        ax.set_axisbelow(True)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        # Slightly larger legend for readability, as requested.
        ax.legend(
            frameon=False, loc="upper center", bbox_to_anchor=(0.5, 1.12),
            ncol=6, fontsize=11, markerscale=1.30, handlelength=2.4,
            columnspacing=1.2
        )

        fig.tight_layout(rect=[0, 0, 1, 0.95])
        save_figure(fig, fig_dir / "06_Intermittent_FDC_Thresholds_Q70_to_Q20.png")
        plt.close(fig)
        gc.collect()
    # 7. Annual completeness heatmap: warm color for excluded, green for usable
    if not annual_ret.empty and not retained.empty:
        a = annual_ret.copy()
        a["Usable_numeric"] = (a["Usable_Year"] == "Yes").astype(float)
        heat = a.pivot(index="Station", columns="Year", values="Usable_numeric")
        order = retained.sort_values("First_Valid_Date")["Station"].tolist()
        heat = heat.reindex([s for s in order if s in heat.index])

        years = heat.columns.to_numpy()
        arr = heat.to_numpy(float)
        from matplotlib.colors import ListedColormap
        cmap = ListedColormap([COLOR_NOT_USABLE, COLOR_USABLE])
        fig, ax = plt.subplots(figsize=(10.0, max(5.5, min(7.5, 0.12 * len(heat) + 1.3))))
        image_obj = ax.imshow(arr, aspect="auto", interpolation="nearest", vmin=0, vmax=1, cmap=cmap)
        ax.set_yticks(np.arange(len(heat.index)))
        ax.set_yticklabels(heat.index)
        if len(years):
            step = max(1, int(math.ceil(len(years) / 12)))
            xt = np.arange(0, len(years), step)
            ax.set_xticks(xt)
            ax.set_xticklabels([str(int(years[i])) for i in xt], rotation=45, ha="right")
        ax.set_xlabel("Calendar year")
        ax.set_ylabel("Station")
        cbar = fig.colorbar(image_obj, ax=ax, pad=0.01, ticks=[0.25, 0.75])
        cbar.ax.set_yticklabels(["Excluded", "Usable"])
        fig.tight_layout()
        save_figure(fig, fig_dir / "07_Annual_Completeness_Heatmap.png")
        plt.close(fig)
        gc.collect()
    # 8. Regime composition
    if not retained.empty:
        counts = retained["Flow_Regime"].value_counts().reindex(
            ["Perennial", "Intermittent"]
        ).fillna(0)
        fig, ax = plt.subplots(figsize=(6.0, 4.5))
        ax.bar(counts.index, counts.values,
               color=[COLOR_PERENNIAL, COLOR_INTERMITTENT],
               edgecolor="white", linewidth=0.4)
        n = float(counts.sum())
        for i, v in enumerate(counts.values):
            pct = 100.0 * float(v) / n if n else 0.0
            ax.text(i, float(v) + 0.2, f"{int(v)} ({pct:.1f}%)",
                    ha="center", va="bottom")
        ax.set_ylabel("Number of retained stations")
        ax.grid(axis="y", linewidth=0.4, alpha=0.22)
        ax.set_axisbelow(True)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        fig.tight_layout()
        save_figure(fig, fig_dir / "08_Flow_Regime_Composition.png")
        plt.close(fig)
        gc.collect()

    # 9-10. Event-characteristic sensitivity across candidate thresholds.
    if not regime_sensitivity.empty:
        plot_event_sensitivity(
            regime_sensitivity, "Perennial", ["Q95", "Q90", "Q85"],
            fig_dir / "09_Perennial_Event_Threshold_Sensitivity.png"
        )
        plot_event_sensitivity(
            regime_sensitivity, "Intermittent", ["Q70", "Q60", "Q50", "Q40", "Q30", "Q20"],
            fig_dir / "10_Intermittent_Event_Threshold_Sensitivity.png"
        )

    # 11-12. Seasonal onset and recovery distributions for complete events.
    if not monthly_timing.empty:
        plot_timing_heatmaps(
            monthly_timing, "Perennial", ["Q95", "Q90", "Q85"],
            fig_dir / "11_Perennial_Event_Timing_Sensitivity.png"
        )
        plot_timing_heatmaps(
            monthly_timing, "Intermittent", ["Q70", "Q60", "Q50", "Q40", "Q30", "Q20"],
            fig_dir / "12_Intermittent_Event_Timing_Sensitivity.png"
        )
# =============================================================================
# MAIN
# =============================================================================
def main() -> None:
    args = parse_args()
    input_dir = args.input
    output_dir = args.output
    fig_dir = output_dir / "FIGURES"

    output_dir.mkdir(parents=True, exist_ok=True)
    fig_dir.mkdir(parents=True, exist_ok=True)

    # Remove legacy zoomed cutoff figure from older runs so it cannot remain
    # in the output folder after this updated script is executed.
    for legacy_stem in [
        "05_ZeroFlow_Percentage_Near_Regime_Threshold",
        "04b_ZeroFlow_Regime_Near_2DayPerYear_Cutoff",
        "04b_ZeroFlow_Regime_Near_1DayPerYear_Cutoff",
    ]:
        for ext in (".png", ".pdf", ".tiff"):
            legacy_file = fig_dir / f"{legacy_stem}{ext}"
            if legacy_file.exists():
                legacy_file.unlink()

    # Remove only legacy Step 1 table CSVs from older runs.
    for legacy_csv in [
        "00_Method_Summary.csv", "01_Station_QC_and_Screening.csv",
        "02_Retained_Stations.csv", "03_Excluded_Stations.csv",
        "04_Annual_Completeness_All_Stations.csv",
        "05_Annual_Completeness_Retained_Stations.csv",
        "06_Perennial_FDC_Thresholds.csv", "07_Intermittent_FDC_Thresholds.csv",
        "08_Missing_Periods_Retained.csv", "09_Continuous_Valid_Blocks_Retained.csv",
        "10_Zero_Flow_Spells_Retained.csv", "11_Isolated_Zero_QC_Log.csv",
        "12_Regime_Borderline_Stations.csv", "13_Processing_Issues.csv",
    ]:
        old_file = output_dir / legacy_csv
        if old_file.exists():
            old_file.unlink()

    if not input_dir.exists():
        raise FileNotFoundError(f"Input folder does not exist: {input_dir}")

    station_files = sorted(input_dir.glob("*.csv"))
    if not station_files:
        raise FileNotFoundError(f"No CSV files found in: {input_dir}")

    meta = load_metadata(args.metadata, input_dir)

    summaries = []
    annual_parts = []
    missing_rows = []
    block_rows = []
    zero_spell_rows = []
    zero_qc_rows = []
    perennial_rows = []
    intermittent_rows = []
    event_sensitivity_rows = []
    event_catalog_rows = []
    pooling_summary_rows = []
    pooling_event_rows = []
    error_rows = []

    for i, path in enumerate(station_files, start=1):
        try:
            result = analyse_station(path, meta)
            if result is None:
                error_rows.append({"File": path.name, "Issue": "No valid discharge observations"})
                continue

            summaries.append(result["summary"])
            annual_parts.append(result["annual"])
            missing_rows.extend(result["missing"])
            block_rows.extend(result["blocks"])
            zero_spell_rows.extend(result["zero_spells"])
            zero_qc_rows.extend(result["zero_qc"])
            if result["perennial"] is not None:
                perennial_rows.append(result["perennial"])
            if result["intermittent"] is not None:
                intermittent_rows.append(result["intermittent"])
            event_sensitivity_rows.extend(result.get("event_sensitivity", []))
            event_catalog_rows.extend(result.get("event_catalog", []))
            pooling_summary_rows.extend(result.get("pooling_summary", []))
            pooling_event_rows.extend(result.get("pooling_events", []))

            # Station hydrograph: one five-panel figure comparing Raw/IT/IC/MA/SPA.
            station_pool_events = pd.DataFrame(result.get("pooling_events", []))
            station_daily = result.get("event_daily", pd.DataFrame())
            s = result["summary"]
            final_q = pd.to_numeric(pd.Series([s.get("Final_Drought_Threshold_m3s", np.nan)]), errors="coerce").iloc[0]
            final_label = str(s.get("Final_Drought_Threshold_Label", "None"))
            if (not station_pool_events.empty and not station_daily.empty and
                    np.isfinite(final_q) and final_q > 0 and
                    s.get("Main_Analysis_Status") == "Retained"):
                plot_pooling_hydrograph_station(
                    station_daily, station_pool_events, s["Station"], s["Flow_Regime"],
                    final_label, float(final_q), fig_dir
                )

            s = result["summary"]
            print(
                f"[{i:02d}/{len(station_files):02d}] {s['Station']}: "
                f"usable={s['Usable_Years']}, status={s['Main_Analysis_Status']}, "
                f"zero={s['Zero_Flow_pct_UsableYears']:.4f}%, regime={s['Flow_Regime']}"
            )
        except Exception as exc:
            error_rows.append({"File": path.name, "Issue": str(exc)})
            print(f"WARNING {path.name}: {exc}")

    if not summaries:
        raise RuntimeError("No station could be analysed successfully.")

    all_df = pd.DataFrame(summaries).sort_values("Station").reset_index(drop=True)
    retained = all_df.loc[all_df["Main_Analysis_Status"] == "Retained"].copy()
    excluded = all_df.loc[all_df["Main_Analysis_Status"] == "Excluded"].copy()
    retained_stations = set(retained["Station"])

    annual_df = pd.concat(annual_parts, ignore_index=True) if annual_parts else pd.DataFrame()
    annual_ret = annual_df.loc[annual_df["Station"].isin(retained_stations)].copy() if not annual_df.empty else annual_df

    missing_df = pd.DataFrame(missing_rows)
    if not missing_df.empty:
        missing_df = missing_df.loc[missing_df["Station"].isin(retained_stations)].copy()

    block_df = pd.DataFrame(block_rows)
    if not block_df.empty:
        block_df = block_df.loc[block_df["Station"].isin(retained_stations)].copy()

    zero_spell_df = pd.DataFrame(zero_spell_rows)
    if not zero_spell_df.empty:
        zero_spell_df = zero_spell_df.loc[zero_spell_df["Station"].isin(retained_stations)].copy()

    zero_qc_df = pd.DataFrame(zero_qc_rows)
    per_df = pd.DataFrame(perennial_rows)
    int_df = pd.DataFrame(intermittent_rows)
    event_sensitivity_df = pd.DataFrame(event_sensitivity_rows)
    event_catalog_df = pd.DataFrame(event_catalog_rows)
    pooling_summary_df = pd.DataFrame(pooling_summary_rows)
    pooling_events_df = pd.DataFrame(pooling_event_rows)
    regime_sensitivity_df = aggregate_regime_sensitivity(event_sensitivity_df)
    monthly_timing_df = monthly_event_timing(event_catalog_df)
    err_df = pd.DataFrame(error_rows)

    # Borderline stations: within +/-0.25 percentage point of threshold
    borderline = retained.loc[
        (retained["Zero_Flow_pct_UsableYears"] - REGIME_ZERO_PCT_CUTOFF).abs() <= 0.25
    ].copy()
    borderline["Distance_from_Regime_Cutoff_pct_points"] = (
        borderline["Zero_Flow_pct_UsableYears"] - REGIME_ZERO_PCT_CUTOFF
    )
    borderline = borderline.sort_values("Zero_Flow_pct_UsableYears")

    regime_counts = retained["Flow_Regime"].value_counts().to_dict()
    method_summary = pd.DataFrame([
        ["Total station files analysed", len(all_df)],
        ["Retained stations", len(retained)],
        ["Excluded stations", len(excluded)],
        ["Perennial retained stations", regime_counts.get("Perennial", 0)],
        ["Intermittent retained stations", regime_counts.get("Intermittent", 0)],
        ["Annual completeness rule", "<=15 missing days per calendar year"],
        ["Minimum usable years per station", MIN_USABLE_YEARS],
        ["Regime zero-flow cutoff (%)", REGIME_ZERO_PCT_CUTOFF],
        ["Missing-data interpolation", "No"],
        ["Threshold calculation data", "Usable years only"],
        ["Drought event definition", "Q < candidate threshold"],
        ["Gap treatment in event analysis", "Missing/unusable days remain NaN and break events"],
        ["Censoring treatment", "Events touching gaps or record boundaries are flagged"],
        ["Event-duration/Qmin summaries", "Complete (uncensored) events only"],
        ["Final perennial drought threshold", "Q90"],
        ["Final intermittent drought threshold", "Q50"],
        ["Pooling methods compared", "Raw, IT, IC, MA, SPA"],
        ["IT critical inter-event time", f"{IT_CRITICAL_TIME_DAYS} days"],
        ["IC settings", f"tc={IC_CRITICAL_TIME_DAYS} days; pc={IC_CRITICAL_VOLUME_RATIO:.2f}"],
        ["MA setting", f"Centered {MA_WINDOW_DAYS}-day moving average"],
        ["Pooling across missing/unusable gaps", "Never"],
        ["Final duration/deficit summaries", "Complete (uncensored) events only"],
        ["Final hydrograph figures", "One Raw/IT/IC/MA/SPA hydrograph per retained station"],
        ["Regime summary figures", "One all-stations summary for Perennial and one for Intermittent"],
    ], columns=["Metric", "Value"])

    # Publication-quality figures.
    make_plots(
        all_df, retained, annual_ret, per_df, int_df,
        regime_sensitivity_df, monthly_timing_df, fig_dir
    )

    # Final station-by-station pooling comparison using Q90 for perennial and
    # Q50 for intermittent stations.
    plot_pooling_comparison_by_station(pooling_summary_df, pooling_events_df, fig_dir)

    # One comprehensive all-stations figure for each regime.
    plot_pooling_summary_all_stations(pooling_summary_df, fig_dir)

    # Consolidated Excel output: one workbook, one sheet per result table.
    excel_tables = [
        ("Method Summary", method_summary, "Method settings and headline counts."),
        ("Station QC", all_df, "QC, completeness, missingness, zero-flow, and regime information for all stations."),
        ("Retained Stations", retained, "Stations retained for the main analysis."),
        ("Excluded Stations", excluded, "Stations excluded because they do not satisfy the minimum usable-year rule."),
        ("Annual All", annual_df, "Annual completeness metrics for all stations."),
        ("Annual Retained", annual_ret, "Annual completeness metrics for retained stations."),
        ("Perennial FDC", per_df.sort_values("Station") if not per_df.empty else per_df, "FDC threshold diagnostics for perennial stations."),
        ("Intermittent FDC", int_df.sort_values("Station") if not int_df.empty else int_df, "Intermittent FDC diagnostics Q70-Q20; final drought analysis uses fixed Q50."),
        ("Threshold Station Sens", event_sensitivity_df.sort_values(["Flow_Regime", "Station", "Threshold_Exceedance_pct"], ascending=[True, True, False]) if not event_sensitivity_df.empty else event_sensitivity_df, "Station-level drought-event sensitivity metrics for every candidate threshold."),
        ("Regime Sensitivity", regime_sensitivity_df.sort_values(["Flow_Regime", "Threshold_Exceedance_pct"], ascending=[True, False]) if not regime_sensitivity_df.empty else regime_sensitivity_df, "Across-station median and IQR sensitivity diagnostics used in Figures 09-10."),
        ("Event Catalogue", event_catalog_df.sort_values(["Flow_Regime", "Station", "Threshold_Exceedance_pct", "Onset_Date"], ascending=[True, True, False, True]) if not event_catalog_df.empty else event_catalog_df, "Event catalogue for every tested threshold, including onset, recovery, duration, Qmin, deficit, zero-flow contribution, and censoring flags."),
        ("Monthly Timing", monthly_timing_df.sort_values(["Flow_Regime", "Threshold_Exceedance_pct", "Timing", "Month"], ascending=[True, False, True, True]) if not monthly_timing_df.empty else monthly_timing_df, "Monthly distribution of complete-event onset and recovery timing for each threshold."),
        ("Final Pooling Summary", pooling_summary_df.sort_values(["Flow_Regime", "Station", "Pooling_Method"]) if not pooling_summary_df.empty else pooling_summary_df, "Station-level comparison of Raw, IT, IC, MA and SPA using fixed Q90 for perennial and Q50 for intermittent stations."),
        ("Final Pooling Events", pooling_events_df.sort_values(["Flow_Regime", "Station", "Pooling_Method", "Onset_Date"]) if not pooling_events_df.empty else pooling_events_df, "Event catalogue for the final fixed-threshold pooling comparison, including duration and physically comparable deficit volume."),
        ("Missing Periods", missing_df, "Internal missing-data periods for retained stations."),
        ("Valid Blocks", block_df, "Continuous valid-data blocks for retained stations."),
        ("Zero Flow Spells", zero_spell_df, "Observed zero-flow spells for retained stations."),
        ("Isolated Zero QC", zero_qc_df, "Log of isolated-zero corrections. Empty means no records in this run."),
        ("Borderline Regime", borderline, "Stations close to the regime-classification cutoff."),
        ("Processing Issues", err_df, "Processing errors/issues. Empty means no issues were recorded."),
    ]
    excel_path = output_dir / EXCEL_FILENAME
    write_excel_workbook(excel_tables, excel_path)

    print("\n================ STEP 1 COMPLETE ================")
    print(f"Stations analysed : {len(all_df)}")
    print(f"Retained          : {len(retained)}")
    print(f"Excluded          : {len(excluded)}")
    print(f"Perennial         : {regime_counts.get('Perennial', 0)}")
    print(f"Intermittent      : {regime_counts.get('Intermittent', 0)}")
    print(f"Regime cutoff     : {REGIME_ZERO_PCT_CUTOFF:.6f}%")
    print(f"Output folder     : {output_dir}")
    print(f"Event rows        : {len(event_catalog_df)}")
    print(f"Sensitivity rows  : {len(event_sensitivity_df)}")
    print(f"Pooling summaries : {len(pooling_summary_df)}")
    print(f"Pooling events    : {len(pooling_events_df)}")
    print(f"Excel workbook    : {excel_path}")
    print("Separate CSV result files are not created.")


if __name__ == "__main__":
    main()